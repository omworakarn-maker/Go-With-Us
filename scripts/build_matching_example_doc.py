from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/GoWithUs_ตัวอย่างการคำนวณความเข้ากันได้.docx")
BLUE = "243F92"
PALE = "EEF3FF"
GRAY = "F5F6F8"
BORDER = "D9D9D9"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def font_run(run, size=11, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, bold=False, align=None, before=0, after=6, size=11):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    font_run(p.add_run(text), size=size, bold=bold)
    return p


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    font_run(p.add_run(text), size=11, bold=True, color=BLUE)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    font_run(run, size=15 if level == 1 else 12, bold=True)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_run(p.add_run(str(header)), size=10, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_border(cell)
            set_cell_margins(cell)
            set_cell_fill(cell, "FFFFFF" if ridx % 2 == 0 else PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            font_run(p.add_run(str(value)), size=10)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(11)
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Arial"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
font_run(title.add_run("ตัวอย่างการคำนวณความเข้ากันได้ระหว่างผู้ใช้กับทริป"), size=20, bold=True)
add_text(doc, "เอกสารนี้อธิบายวิธีคำนวณคะแนน Smart Matching ของ GoWithUs ด้วยตัวอย่างเดียวตั้งแต่ข้อมูลนำเข้า คะแนนรายด้าน การถ่วงน้ำหนัก จนถึงคะแนนรวม เพื่อให้นำไปอธิบายหรือจัดทำสื่อ Canva ได้ง่าย", align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

add_heading(doc, "สัดส่วนคะแนน", 1)
add_table(doc, ["ปัจจัย", "น้ำหนัก"], [
    ("ความสนใจและหมวดหมู่ทริป", "35%"),
    ("งบประมาณ", "30%"),
    ("จำนวนกิจกรรมต่อวัน", "20%"),
    ("ช่วงเวลาที่ชอบทำกิจกรรม", "15%"),
    ("รวม", "100%"),
], [5.6, 1.2])
add_formula(doc, "คะแนนรวม = (ความสนใจ × 0.35) + (งบประมาณ × 0.30) + (กิจกรรม × 0.20) + (ช่วงเวลา × 0.15)")

add_heading(doc, "ข้อมูลตัวอย่าง", 1)
add_table(doc, ["ปัจจัย", "ผู้ใช้", "ทริป"], [
    ("ความสนใจหรือหมวดหมู่", "ทะเล คาเฟ่ ไหว้พระ", "ทะเล คาเฟ่"),
    ("งบประมาณต่อคน", "1,000 บาท", "1,000 บาท"),
    ("กิจกรรมต่อวัน", "3–4 กิจกรรม ระดับ 5", "3–4 กิจกรรม ระดับ 5"),
    ("ช่วงเวลา", "เช้า เย็น", "เย็น กลางคืน"),
], [2.0, 2.4, 2.4])

doc.add_page_break()
add_heading(doc, "1 ความสนใจและหมวดหมู่ทริป", 1)
add_text(doc, "น้ำหนัก 35% ระบบตรวจว่าความสนใจของผู้ใช้ตรงกับหมวดหมู่หลักและหมวดหมู่เพิ่มเติมของทริปมากน้อยเพียงใด โดยแปลงแต่ละหมวดหมู่เป็นเลข 1 เมื่อเลือก และ 0 เมื่อไม่เลือก แล้วใช้ Cosine Similarity")
add_table(doc, ["หมวดหมู่", "ผู้ใช้", "ทริป"], [("ทะเล", 1, 1), ("คาเฟ่", 1, 1), ("ไหว้พระ", 1, 0)], [4.0, 1.4, 1.4])
add_text(doc, "เวกเตอร์ผู้ใช้ = [1, 1, 1]   เวกเตอร์ทริป = [1, 1, 0]", bold=True)
add_text(doc, "ผลคูณจุด = (1×1) + (1×1) + (1×0) = 2")
add_text(doc, "ความยาวเวกเตอร์ผู้ใช้ = √3 และความยาวเวกเตอร์ทริป = √2")
add_formula(doc, "คะแนนความสนใจ = 2 ÷ (√3 × √2) × 100 = 81.65% ≈ 82%")
add_formula(doc, "คะแนนหลังถ่วงน้ำหนัก = 82 × 0.35 = 28.70 คะแนน")

add_heading(doc, "2 งบประมาณ", 1)
add_text(doc, "น้ำหนัก 30% ระบบเปรียบเทียบงบประมาณที่ผู้ใช้รับได้กับงบประมาณต่อคนของทริป")
add_table(doc, ["เงื่อนไข", "คะแนนงบประมาณ"], [
    ("งบผู้ใช้มากกว่าหรือเท่ากับงบทริป", "100%"),
    ("งบผู้ใช้น้อยกว่า 50% ของงบทริป", "0%"),
    ("งบผู้ใช้อยู่ระหว่าง 50% ถึงต่ำกว่างบทริป", "งบผู้ใช้ ÷ งบทริป × 100"),
], [4.8, 2.0])
add_text(doc, "ตัวอย่างนี้ ผู้ใช้มีงบ 1,000 บาท และทริปมีงบ 1,000 บาท งบจึงเพียงพอ")
add_formula(doc, "คะแนนงบประมาณ = 100%")
add_formula(doc, "คะแนนหลังถ่วงน้ำหนัก = 100 × 0.30 = 30.00 คะแนน")
add_text(doc, "ตัวอย่างเพิ่มเติม หากผู้ใช้มีงบ 900 บาท และทริปมีงบ 1,000 บาท คะแนนงบประมาณเท่ากับ 900 ÷ 1,000 × 100 = 90%")

doc.add_page_break()
add_heading(doc, "3 จำนวนกิจกรรมต่อวัน", 1)
add_text(doc, "น้ำหนัก 20% ระบบหาค่าเฉลี่ยจำนวนกิจกรรมต่อวันของทริป แล้วแปลงเป็นระดับเดียวกับคำตอบของผู้ใช้")
add_table(doc, ["จำนวนกิจกรรมเฉลี่ยต่อวัน", "รหัสระดับ"], [("1–2 กิจกรรม", 2), ("3–4 กิจกรรม", 5), ("5 กิจกรรมขึ้นไป", 8)], [5.2, 1.6])
add_text(doc, "ตัวอย่างนี้ ผู้ใช้และทริปอยู่ระดับ 5 เหมือนกัน ระยะห่างจึงเป็นศูนย์")
add_formula(doc, "คะแนนกิจกรรม = [1 − (|ระดับผู้ใช้ − ระดับทริป| ÷ 6)] × 100")
add_formula(doc, "= [1 − (|5 − 5| ÷ 6)] × 100 = 100%")
add_formula(doc, "คะแนนหลังถ่วงน้ำหนัก = 100 × 0.20 = 20.00 คะแนน")
add_table(doc, ["ระดับผู้ใช้", "ระดับทริป", "ผลต่าง", "คะแนน"], [(2, 2, 0, "100%"), (2, 5, 3, "50%"), (2, 8, 6, "0%"), (5, 8, 3, "50%")], [1.7, 1.7, 1.7, 1.7])
add_text(doc, "หมายเหตุ ค่า 2 5 และ 8 เป็นรหัสแทนระดับที่โครงการกำหนด ไม่ใช่จำนวนกิจกรรมจริงและไม่ใช่ค่ามาตรฐานที่งานวิจัยกำหนด", bold=True)

add_heading(doc, "4 ช่วงเวลาที่ชอบทำกิจกรรม", 1)
add_text(doc, "น้ำหนัก 15% ระบบใช้ช่วงเวลา 4 ช่อง ได้แก่ เช้า กลางวัน เย็น และกลางคืน จากนั้นแปลงเป็น 0 และ 1 แล้วคำนวณ Cosine Similarity")
add_table(doc, ["ช่วงเวลา", "ผู้ใช้", "ทริป"], [("เช้า", 1, 0), ("กลางวัน", 0, 0), ("เย็น", 1, 1), ("กลางคืน", 0, 1)], [4.0, 1.4, 1.4])
add_text(doc, "เวกเตอร์ผู้ใช้ = [1, 0, 1, 0]   เวกเตอร์ทริป = [0, 0, 1, 1]", bold=True)
add_text(doc, "ผลคูณจุด = 1 ความยาวของเวกเตอร์ทั้งสองเท่ากับ √2")
add_formula(doc, "คะแนนช่วงเวลา = 1 ÷ (√2 × √2) × 100 = 50%")
add_formula(doc, "คะแนนหลังถ่วงน้ำหนัก = 50 × 0.15 = 7.50 คะแนน")

doc.add_page_break()
add_heading(doc, "สรุปคะแนน", 1)
add_table(doc, ["ปัจจัย", "คะแนนย่อย", "น้ำหนัก", "คะแนนหลังถ่วงน้ำหนัก"], [
    ("ความสนใจ", "82%", "35%", "28.70"),
    ("งบประมาณ", "100%", "30%", "30.00"),
    ("จำนวนกิจกรรม", "100%", "20%", "20.00"),
    ("ช่วงเวลา", "50%", "15%", "7.50"),
    ("รวม", "", "100%", "86.20"),
], [2.4, 1.3, 1.3, 1.8])
add_formula(doc, "คะแนนรวม = 28.70 + 30.00 + 20.00 + 7.50 = 86.20%")
add_text(doc, "ระบบปัดคะแนนเป็นจำนวนเต็ม จึงแสดงผลเป็น 86% Match", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=12, size=16)

add_heading(doc, "เงื่อนไขเพิ่มเติมของระบบ", 1)
for item in [
    "หากทริปเต็มแล้ว คะแนนรวมเป็น 0%",
    "หากทริปสิ้นสุดไปแล้ว คะแนนรวมเป็น 0%",
    "หากงบทริปสูงกว่างบผู้ใช้มากกว่า 2 เท่า คะแนนรวมถูกจำกัดไว้ไม่เกิน 39%",
    "หากปัจจัยบางด้านไม่มีข้อมูล ระบบตัดปัจจัยนั้นออกและหารด้วยผลรวมน้ำหนักของข้อมูลที่เหลือ",
    "คะแนนรายด้านและคะแนนรวมถูกปัดเป็นจำนวนเต็มก่อนนำไปแสดงผล",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    font_run(p.add_run(item), size=11)

add_heading(doc, "ข้อความสรุปสำหรับนำเสนอ", 1)
add_text(doc, "ระบบ Smart Matching ของ GoWithUs เปรียบเทียบผู้ใช้กับทริปจากความสนใจ 35% งบประมาณ 30% จำนวนกิจกรรมต่อวัน 20% และช่วงเวลา 15% ความสนใจและช่วงเวลาใช้ Cosine Similarity ส่วนงบประมาณและระดับกิจกรรมใช้การเปรียบเทียบความแตกต่างตามกฎของระบบ จากตัวอย่าง ผู้ใช้กับทริปได้คะแนนย่อย 82% 100% 100% และ 50% ตามลำดับ เมื่อนำมาถ่วงน้ำหนักจึงได้คะแนนความเข้ากันได้รวม 86%")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
