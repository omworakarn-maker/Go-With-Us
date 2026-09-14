from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "GoWithUs_แบบสอบถามความสำคัญของปัจจัย_ฉบับง่าย.docx"


def format_run(run, size=14, bold=False, color=None):
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def set_text(cell, value, bold=False, color=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(value)
    format_run(r, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def heading(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(value)
    format_run(r, size=17, bold=True, color="24449A")


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2)
sec.right_margin = Cm(2)
normal = doc.styles["Normal"]
normal.font.name = "TH Sarabun New"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
normal.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("แบบสอบถามปัจจัยที่มีผลต่อการเลือกทริป")
format_run(r, size=20, bold=True, color="24449A")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("เพื่อกำหนดน้ำหนักในระบบ Smart Matching ของ GoWithUs")
format_run(r, size=16, bold=True)

heading(doc, "คำชี้แจง")
doc.add_paragraph(
    "แบบสอบถามนี้จัดทำขึ้นเพื่อศึกษาความสำคัญของปัจจัยที่มีผลต่อการเลือกทริป ได้แก่ ความสนใจ "
    "งบประมาณ จำนวนกิจกรรมต่อวัน และช่วงเวลา เพื่อนำผลไปกำหนดค่าน้ำหนักในระบบ Smart Matching "
    "ของแอปพลิเคชัน GoWithUs โดยวิเคราะห์จากคะแนนเฉลี่ยของกลุ่มตัวอย่าง ใช้เวลาประมาณ 1–2 นาที "
    "ไม่มีคำตอบถูกหรือผิด และรายงานผลโดยรวมโดยไม่ระบุตัวตน"
)

heading(doc, "ส่วนที่ 1 ข้อมูลเบื้องต้น")
doc.add_paragraph("1. คุณเคยเดินทางท่องเที่ยวร่วมกับเพื่อนหรือบุคคลอื่นหรือไม่?")
doc.add_paragraph("☐ เคย     ☐ ไม่เคย")
doc.add_paragraph("2. คุณสนใจใช้แอปพลิเคชันสำหรับค้นหาทริปหรือเพื่อนร่วมเดินทางหรือไม่?")
doc.add_paragraph("☐ สนใจ     ☐ ไม่แน่ใจ     ☐ ไม่สนใจ")

heading(doc, "ส่วนที่ 2 ให้คะแนนความสำคัญ")
doc.add_paragraph(
    "เมื่อคุณกำลังเลือกทริป ให้คะแนนความสำคัญของแต่ละปัจจัยตั้งแต่ 1–5 "
    "โดย 1 หมายถึงไม่สำคัญเลย และ 5 หมายถึงสำคัญมากที่สุด"
)

factors = [
    ("ความสนใจ", "ประเภททริปตรงกับสิ่งที่คุณชอบ เช่น ทะเล ภูเขา คาเฟ่ หรือไหว้พระ"),
    ("งบประมาณ", "ค่าใช้จ่ายเฉลี่ยต่อคนอยู่ในระดับที่คุณสามารถจ่ายได้"),
    ("จำนวนกิจกรรมต่อวัน", "จำนวนกิจกรรมและความแน่นของตารางตรงกับรูปแบบการเที่ยวที่คุณต้องการ"),
    ("ช่วงเวลา", "ช่วงเช้า กลางวัน เย็น หรือกลางคืนตรงกับช่วงที่คุณสะดวกหรือชอบทำกิจกรรม"),
]

table = doc.add_table(rows=1, cols=7)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["ปัจจัย", "คำอธิบาย", "1", "2", "3", "4", "5"]
for i, value in enumerate(headers):
    set_text(table.rows[0].cells[i], value, bold=True, color="FFFFFF", center=True)
    shade(table.rows[0].cells[i], "24449A")
for factor, description in factors:
    cells = table.add_row().cells
    set_text(cells[0], factor, bold=True)
    set_text(cells[1], description)
    for i in range(2, 7):
        set_text(cells[i], "☐", center=True)

doc.add_paragraph("ความหมายของคะแนน: 1 = ไม่สำคัญเลย, 2 = สำคัญเล็กน้อย, 3 = ปานกลาง, 4 = สำคัญ, 5 = สำคัญมากที่สุด")

heading(doc, "ส่วนที่ 3 เลือกปัจจัยที่สำคัญที่สุด")
doc.add_paragraph("หากเลือกได้เพียงหนึ่งข้อ ปัจจัยใดมีผลต่อการตัดสินใจเลือกทริปของคุณมากที่สุด?")
doc.add_paragraph("☐ ความสนใจ")
doc.add_paragraph("☐ งบประมาณ")
doc.add_paragraph("☐ จำนวนกิจกรรมต่อวัน")
doc.add_paragraph("☐ ช่วงเวลา")

heading(doc, "ส่วนที่ 4 ปัจจัยเพิ่มเติม (ไม่บังคับ)")
doc.add_paragraph("มีปัจจัยอื่นใดที่คุณใช้พิจารณาเลือกทริปหรือไม่ เช่น ความปลอดภัย สถานที่ หรือจำนวนผู้ร่วมทริป?")
doc.add_paragraph("............................................................................................................................")
doc.add_paragraph("............................................................................................................................")

heading(doc, "วิธีคำนวณสำหรับผู้วิจัย")
doc.add_paragraph(
    "คำนวณค่าเฉลี่ยคะแนนของแต่ละปัจจัย แล้วแปลงเป็นน้ำหนักรวม 100% ด้วยสูตร: "
    "น้ำหนักปัจจัย (%) = ค่าเฉลี่ยของปัจจัย ÷ ผลรวมค่าเฉลี่ยทั้ง 4 ปัจจัย × 100 "
    "พร้อมรายงานจำนวนผู้ตอบ ค่าเฉลี่ย และส่วนเบี่ยงเบนมาตรฐาน ห้ามปรับคำตอบเพื่อให้ได้ค่าน้ำหนักเดิม"
)
doc.add_paragraph(
    "หมายเหตุ: แบบสอบถามฉบับนี้เป็นการให้คะแนนโดยตรง (direct rating) ซึ่งตอบง่ายกว่า AHP "
    "ควรระบุวิธีดังกล่าวตามจริงในระเบียบวิธีวิจัย ไม่ควรเรียกว่า AHP"
)

doc.save(OUTPUT)
print(OUTPUT)
