from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path('/Users/worakanp/Desktop/Go-with-us-1/reports/รายงานฉบับเต็ม_ปรับแก้.docx')
OUTPUT = Path('/Users/worakanp/Desktop/Go-with-us-1/reports/รายงานฉบับเต็ม_ปรับแก้_สารบรรณ16.docx')
FONT = 'TH Sarabun New'
SIZE = '32'  # OOXML stores point sizes in half-points.


def child(parent, tag):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        if tag in ('w:rPr', 'w:rFonts') and parent.tag == qn('w:r'):
            parent.insert(0, element)
        elif tag == 'w:rFonts':
            parent.insert(0, element)
        else:
            parent.append(element)
    return element


def set_font(properties):
    fonts = child(properties, 'w:rFonts')
    for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{script}'), FONT)
    # Remove theme values so they cannot override the explicit typeface.
    for script in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn(f'w:{script}'), None)
    for tag in ('w:sz', 'w:szCs'):
        child(properties, tag).set(qn('w:val'), SIZE)


doc = Document(SOURCE)
toc_pages = {
    'บทที่ 1 บทนำ': 6,
    'บทที่ 2 ทฤษฎีและงานวิจัยที่เกี่ยวข้อง': 15,
    'บทที่ 3 วิธีการดำเนินการ': 29,
    'บทที่ 4 ผลการดำเนินงานและผลการทดสอบระบบ': 52,
    'บทที่ 5 สรุปผล อภิปรายผล และข้อเสนอแนะ': 132,
    'อ้างอิง': 136,
}
for paragraph in doc.paragraphs:
    label = paragraph.text.split('\t', 1)[0]
    if label in toc_pages and '\t' in paragraph.text:
        paragraph.clear()
        paragraph.add_run(f'{label}\t{toc_pages[label]}')

parts = [doc.part.element, doc.styles.element]
for section in doc.sections:
    parts.extend((section.header._element, section.footer._element))

seen = set()
for root in parts:
    if id(root) in seen:
        continue
    seen.add(id(root))
    for run in root.xpath('.//w:r'):
        set_font(child(run, 'w:rPr'))
    for style in root.xpath('.//w:style'):
        set_font(child(style, 'w:rPr'))
    if root.tag == qn('w:styles'):
        defaults = root.find(qn('w:docDefaults'))
        if defaults is not None:
            run_defaults = child(defaults, 'w:rPrDefault')
            set_font(child(run_defaults, 'w:rPr'))

doc.save(OUTPUT)
print(OUTPUT)
