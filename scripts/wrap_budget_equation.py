from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
path='/tmp/report_polished/รายงานฉบับเต็ม_จัดหน้าพร้อมส่ง.docx'
d=Document(path)
p=next(p for p in d.paragraphs if 'เมื่อ' in [n.text for n in p._p.xpath('.//m:t')])
m=p._p.xpath('.//m:oMath')[0]
for child in list(m)[4:]:m.remove(child)
el=deepcopy(p._p)
p._p.addnext(el)
condition=Paragraph(el,p._parent)
condition.clear()
r=condition.add_run('เมื่อ 0 ≤ UserBudget < TripBudget')
r.font.name='TH Sarabun New';r.font.size=Pt(16)
condition.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
condition.paragraph_format.space_before=Pt(0)
condition.paragraph_format.space_after=Pt(6)
condition.paragraph_format.keep_with_next=False
d.save(path)
