from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

pdf=next(Path(sys.argv[1] if len(sys.argv)>1 else '/tmp/report_polished_render2').glob('*.pdf'))
texts=[p.extract_text() for p in PdfReader(pdf).pages]
path=Path(sys.argv[2] if len(sys.argv)>2 else '/tmp/report_polished/รายงานฉบับเต็ม_จัดหน้าพร้อมส่ง.docx')
font_name=sys.argv[3] if len(sys.argv)>3 else 'TH Sarabun New'
d=Document(path)
has_continued=any(p.text=='สารบัญ (ต่อ)' for p in d.paragraphs)
mapping={}
for p in d.paragraphs:
    if not p.style.name.startswith('toc'):continue
    label=p.text.split('\t')[0]
    m=re.match(r'^(\d+\.\d+)\s',label)
    chapter=re.match(r'^บทที่\s+(\d)',label)
    key=m.group(1) if m else chapter.group(1)+'.1' if chapter else None
    if key:
        key_pattern=r'\s*\.\s*'.join(re.escape(k) for k in key.split('.'))
        found=[i+1 for i,t in enumerate(texts) if i>4 and re.search(r'(?m)^\s*'+key_pattern+r'(?:\s+|[^0-9.])(?!\s*\.)',t)]
        if not found:raise ValueError(key)
        page=found[0]
    else:
        found=[i+1 for i,t in enumerate(texts) if i>4 and t.lstrip().startswith('อ') and '[1]' in t[:350]]
        if not found:raise ValueError('references')
        page=found[0]
    mapping[label]=page
    p.clear();r=p.add_run(label+'\t'+str(page));r.font.name=font_name;r.font.size=Pt(16)
    from docx.oxml.ns import qn
    for a in ('ascii','hAnsi','eastAsia','cs'):r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+a),font_name)
    if label.startswith('บทที่ 3 ') and not has_continued:
        p.paragraph_format.page_break_before=False
        h=p.insert_paragraph_before('สารบัญ (ต่อ)')
        h.paragraph_format.page_break_before=True
        h.paragraph_format.keep_with_next=True
        h.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.space_after=Pt(8)
        h.runs[0].font.name=font_name;h.runs[0].font.size=Pt(16);h.runs[0].bold=True
d.save(path)
print('Pages',len(texts))
print(mapping)
