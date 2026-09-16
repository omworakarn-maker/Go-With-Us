from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path


OUT = "/Users/worakanp/Desktop/Go-with-us-1/reports/บทที่3_ฉบับแก้ไขทดลอง.docx"
DIAGRAM_DIR = Path("/tmp/go-with-us-chapter3-diagrams")
DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
THAI_FONT = "/System/Library/Fonts/Supplemental/Tahoma.ttf"


def diagram_font(size, bold=False):
    path = "/System/Library/Fonts/Supplemental/Tahoma Bold.ttf" if bold else THAI_FONT
    return ImageFont.truetype(path, size)


def centered_text(draw, box, text, font, fill="#152238"):
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + (len(lines) - 1) * 8
    y = y1 + (y2 - y1 - total) / 2
    for line, h in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=font, fill=fill)
        y += h + 8


def rounded_box(draw, box, text, fill="#EDF4FB", outline="#245780", width=4, font_size=30):
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=width)
    centered_text(draw, box, text, diagram_font(font_size, True))


def arrow(draw, start, end, fill="#245780", width=6):
    draw.line([start, end], fill=fill, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    pts = [end,
           (end[0] - size * math.cos(angle - 0.5), end[1] - size * math.sin(angle - 0.5)),
           (end[0] - size * math.cos(angle + 0.5), end[1] - size * math.sin(angle + 0.5))]
    draw.polygon(pts, fill=fill)


def save_architecture_diagram():
    path = DIAGRAM_DIR / "architecture.png"
    im = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(im)
    rounded_box(d, (70, 100, 380, 260), "Native iOS\nSwiftUI", fill="#E8F1FF")
    rounded_box(d, (70, 450, 380, 610), "Admin Backoffice\nReact และ Vite", fill="#E8F1FF")
    rounded_box(d, (570, 275, 930, 435), "Backend REST API\nNode.js Express Prisma", fill="#DDEBFA")
    rounded_box(d, (1120, 100, 1430, 260), "PostgreSQL\nSupabase", fill="#EAF7F0", outline="#2B7A58")
    rounded_box(d, (1120, 450, 1430, 610), "บริการภายนอก\nอีเมล Gemini แผนที่", fill="#FFF4DF", outline="#9A6718")
    arrow(d, (380, 180), (570, 330)); arrow(d, (380, 530), (570, 380))
    arrow(d, (930, 330), (1120, 180)); arrow(d, (930, 380), (1120, 530))
    d.text((455, 190), "HTTPS JSON และ JWT", font=diagram_font(25), fill="#4A5568")
    d.text((985, 200), "Prisma", font=diagram_font(25), fill="#4A5568")
    d.text((965, 500), "API", font=diagram_font(25), fill="#4A5568")
    im.save(path, dpi=(180, 180)); return path


def save_use_case_diagram():
    path = DIAGRAM_DIR / "use-case.png"
    im = Image.new("RGB", (1500, 920), "white")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((380, 60, 1120, 860), radius=28, fill="#F7FAFC", outline="#245780", width=5)
    d.text((590, 82), "ระบบ Go With Us", font=diagram_font(34, True), fill="#152238")
    rounded_box(d, (40, 180, 300, 310), "ผู้ใช้งาน", fill="#E8F1FF")
    rounded_box(d, (40, 600, 300, 730), "ผู้ดูแลระบบ", fill="#FFF0F0", outline="#B94A48")
    cases=[("สมัครและยืนยัน OTP",170),("ตอบแบบสอบถามและแก้ไขโปรไฟล์",285),("ค้นหา สร้าง และเข้าร่วมทริป",400),("สนทนา แจ้งเตือน และรายงาน",515),("จัดการผู้ใช้ ทริป รายงาน และประกาศ",685)]
    for text_value,y in cases:
        rounded_box(d,(500,y,1000,y+78),text_value,fill="white",font_size=25)
    for y in (209,324,439,554): arrow(d,(300,245),(500,y),width=4)
    arrow(d,(300,665),(500,724),fill="#B94A48",width=4)
    im.save(path,dpi=(180,180)); return path


def save_otp_sequence_diagram():
    path = DIAGRAM_DIR / "otp-sequence.png"
    im = Image.new("RGB", (1500, 980), "white")
    d = ImageDraw.Draw(im)
    xs=[150,520,900,1270]; labels=["ผู้ใช้","Native iOS","Backend","บริการอีเมล"]
    for x,label in zip(xs,labels):
        rounded_box(d,(x-120,45,x+120,125),label,fill="#E8F1FF",font_size=25)
        d.line((x,125,x,920),fill="#AAB4C0",width=3)
    steps=[(150,520,200,"กรอกข้อมูลสมัคร"),(520,900,310,"POST ลงทะเบียน"),(900,1270,420,"ส่งรหัส OTP"),(150,520,550,"กรอก OTP"),(520,900,660,"ตรวจสอบ OTP"),(900,520,790,"สร้าง User เมื่อตรวจสอบสำเร็จ")]
    for a,b,y,label in steps:
        arrow(d,(a,y),(b,y),width=5)
        mid=(a+b)//2; box=d.textbbox((0,0),label,font=diagram_font(24));
        d.rectangle((mid-(box[2]-box[0])/2-8,y-40,mid+(box[2]-box[0])/2+8,y-8),fill="white")
        d.text((mid-(box[2]-box[0])/2,y-39),label,font=diagram_font(24),fill="#152238")
    d.text((760,855),"หากรหัสผิดหรือหมดอายุ ระบบยังไม่สร้างบัญชีผู้ใช้",font=diagram_font(23),fill="#B94A48")
    im.save(path,dpi=(180,180)); return path


def save_er_diagram():
    path = DIAGRAM_DIR / "er-diagram.png"
    im = Image.new("RGB", (1600, 1050), "white")
    d = ImageDraw.Draw(im)
    boxes={
        "users":(650,60,950,180), "trips":(650,420,950,540), "participants":(110,420,430,540),
        "messages":(1170,420,1490,540), "notifications":(110,760,430,880), "user_reports":(650,760,950,880),
        "pending_registrations":(1080,60,1530,180)}
    labels={"users":"users\nPK id","trips":"trips\nPK id  FK creator_id","participants":"participants\nFK user_id  FK trip_id","messages":"messages\nFK sender_id  trip_id receiver_id","notifications":"notifications\nFK user_id อาจว่าง","user_reports":"user_reports\nFK reporter_id reported_id","pending_registrations":"pending_registrations\nemail และ OTP ชั่วคราว"}
    for key,box in boxes.items(): rounded_box(d,box,labels[key],fill="#F3F7FA",font_size=23)
    arrow(d,(800,180),(800,420),width=4); d.text((815,280),"1 สร้าง N",font=diagram_font(22),fill="#4A5568")
    arrow(d,(650,480),(430,480),width=4); arrow(d,(950,480),(1170,480),width=4)
    arrow(d,(720,180),(320,420),width=4); arrow(d,(880,180),(1330,420),width=4)
    arrow(d,(720,180),(300,760),width=4); arrow(d,(800,180),(800,760),width=4)
    d.line((950,120,1080,120),fill="#9A6718",width=4)
    d.text((965,78),"ยืนยันสำเร็จแล้วจึงสร้าง",font=diagram_font(20),fill="#9A6718")
    d.text((460,455),"N : 1",font=diagram_font(21),fill="#4A5568")
    d.text((1000,455),"1 : N",font=diagram_font(21),fill="#4A5568")
    im.save(path,dpi=(180,180)); return path


def save_matching_flowchart():
    path = DIAGRAM_DIR / "matching-flowchart.png"
    im = Image.new("RGB", (1500, 1500), "white")
    d = ImageDraw.Draw(im)
    rounded_box(d, (510, 35, 990, 145), "รับคำขอหน้าแนะนำทริปพร้อม JWT", fill="#E8F1FF", font_size=27)
    arrow(d, (750, 145), (750, 235))
    rounded_box(d, (510, 235, 990, 345), "อ่านข้อมูลผู้ใช้และรายการทริป", fill="#F3F7FA", font_size=27)
    arrow(d, (750, 345), (750, 435))
    rounded_box(d, (510, 435, 990, 555), "ทริปยังไม่สิ้นสุดและจำนวนผู้เข้าร่วมยังไม่เต็ม", fill="#FFF4DF", outline="#9A6718", font_size=25)
    arrow(d, (750, 555), (750, 650))
    rounded_box(d, (80, 650, 390, 790), "ความสนใจ\nCosine Similarity", fill="#EAF7F0", outline="#2B7A58", font_size=24)
    rounded_box(d, (415, 650, 725, 790), "งบประมาณ\nคะแนนเชิงสัดส่วน", fill="#EAF7F0", outline="#2B7A58", font_size=24)
    rounded_box(d, (750, 650, 1060, 790), "จำนวนสถานที่ต่อวัน\nคะแนนจากระยะห่าง", fill="#EAF7F0", outline="#2B7A58", font_size=24)
    rounded_box(d, (1085, 650, 1395, 790), "ช่วงเวลา\nCosine Similarity", fill="#EAF7F0", outline="#2B7A58", font_size=24)
    for x in (235, 570, 905, 1240): arrow(d, (x, 790), (750, 930), width=4)
    rounded_box(d, (470, 930, 1030, 1050), "รวมเฉพาะคะแนนย่อยที่คำนวณได้", fill="#DDEBFA", font_size=27)
    arrow(d, (750, 1050), (750, 1140))
    rounded_box(d, (470, 1140, 1030, 1260), "คำนวณค่าเฉลี่ยและปัดเป็นจำนวนเต็ม", fill="#DDEBFA", font_size=27)
    arrow(d, (750, 1260), (750, 1350))
    rounded_box(d, (470, 1350, 1030, 1460), "เรียงคะแนนจากมากไปน้อยและส่งผลเป็น JSON", fill="#E8F1FF", font_size=25)
    im.save(path, dpi=(180, 180)); return path


def save_join_sequence_diagram():
    path = DIAGRAM_DIR / "join-sequence.png"
    im = Image.new("RGB", (1600, 1250), "white")
    d = ImageDraw.Draw(im)
    xs = [140, 500, 850, 1210, 1480]
    labels = ["ผู้ใช้", "Native iOS", "Backend", "ฐานข้อมูล", "Notification"]
    for x, label in zip(xs, labels):
        rounded_box(d, (x - 105, 35, x + 105, 120), label, fill="#E8F1FF", font_size=23)
        d.line((x, 120, x, 1170), fill="#AAB4C0", width=3)
    steps = [
        (140, 500, 190, "กดยืนยันเข้าร่วม"),
        (500, 850, 300, "POST /trips/:id/join พร้อม JWT"),
        (850, 1210, 410, "ค้นหาทริปและผู้เข้าร่วม"),
        (1210, 850, 520, "ส่งข้อมูลทริปกลับ"),
        (850, 1210, 650, "สร้างหรืออัปเดต Participant"),
        (850, 1480, 790, "สร้างการแจ้งเตือนให้สมาชิกเดิม"),
        (850, 500, 930, "ส่งสถานะสำเร็จและข้อมูล Participant"),
        (500, 140, 1050, "แสดงสถานะเข้าร่วมแล้ว"),
    ]
    for a, b, y, label in steps:
        arrow(d, (a, y), (b, y), width=5)
        bbox = d.textbbox((0, 0), label, font=diagram_font(22))
        mid = (a + b) / 2
        d.rectangle((mid - (bbox[2]-bbox[0])/2 - 7, y-37, mid + (bbox[2]-bbox[0])/2 + 7, y-7), fill="white")
        d.text((mid - (bbox[2]-bbox[0])/2, y-36), label, font=diagram_font(22), fill="#152238")
    d.text((750, 575), "ถ้าทริปเต็ม ระบบส่งข้อผิดพลาดและไม่สร้างข้อมูล", font=diagram_font(21), fill="#B94A48")
    im.save(path, dpi=(180, 180)); return path


def save_create_trip_sequence_diagram():
    path = DIAGRAM_DIR / "create-trip-sequence.png"
    im = Image.new("RGB", (1600, 1220), "white")
    d = ImageDraw.Draw(im)
    xs = [150, 500, 860, 1220, 1480]
    labels = ["ผู้ใช้", "Native iOS", "Backend", "Gemini API", "ฐานข้อมูล"]
    for x, label in zip(xs, labels):
        rounded_box(d, (x - 105, 35, x + 105, 120), label, fill="#E8F1FF", font_size=23)
        d.line((x, 120, x, 1140), fill="#AAB4C0", width=3)
    steps = [
        (150, 500, 185, "กรอกข้อมูลและกดสร้างทริป"),
        (500, 860, 300, "POST /trips พร้อม JWT และข้อมูลทริป"),
        (860, 1480, 410, "ตรวจสอบผู้สร้างจาก users"),
        (860, 1220, 540, "ขอสร้าง embedding แบบไม่บังคับ"),
        (860, 1480, 690, "สร้าง trips และ Participant ของผู้สร้าง"),
        (1480, 860, 820, "ส่งข้อมูลทริปที่สร้างแล้ว"),
        (860, 500, 950, "HTTP 201 พร้อมข้อมูล Trip"),
        (500, 150, 1070, "แสดงทริปที่สร้างสำเร็จ"),
    ]
    for a, b, y, label in steps:
        arrow(d, (a, y), (b, y), width=5)
        bbox = d.textbbox((0, 0), label, font=diagram_font(21))
        mid = (a + b) / 2
        d.rectangle((mid-(bbox[2]-bbox[0])/2-7, y-36, mid+(bbox[2]-bbox[0])/2+7, y-7), fill="white")
        d.text((mid-(bbox[2]-bbox[0])/2, y-35), label, font=diagram_font(21), fill="#152238")
    d.text((900, 600), "หาก Gemini ใช้งานไม่ได้ ระบบยังสร้างทริปได้", font=diagram_font(20), fill="#9A6718")
    im.save(path, dpi=(180, 180)); return path


def save_message_sequence_diagram():
    path = DIAGRAM_DIR / "message-sequence.png"
    im = Image.new("RGB", (1600, 1300), "white")
    d = ImageDraw.Draw(im)
    xs = [140, 480, 820, 1160, 1480]
    labels = ["ผู้ส่ง", "Native iOS", "Backend", "ฐานข้อมูล", "สมาชิกคนอื่น"]
    for x, label in zip(xs, labels):
        rounded_box(d, (x - 105, 35, x + 105, 120), label, fill="#E8F1FF", font_size=22)
        d.line((x, 120, x, 1210), fill="#AAB4C0", width=3)
    steps = [
        (140, 480, 180, "พิมพ์ข้อความหรือเลือกรูปภาพ"),
        (480, 820, 290, "ส่งข้อความพร้อม tripId และ JWT"),
        (820, 1160, 405, "ตรวจสอบสถานะ Participant"),
        (1160, 820, 520, "ยืนยันว่าเป็นสมาชิกทริป"),
        (820, 1160, 650, "บันทึก Message"),
        (1160, 820, 770, "ส่งข้อมูลข้อความกลับ"),
        (820, 1480, 900, "WebSocket และ Push Notification"),
        (820, 480, 1030, "HTTP 201 พร้อม Message"),
        (480, 140, 1140, "แสดงข้อความในห้องสนทนา"),
    ]
    for a, b, y, label in steps:
        arrow(d, (a, y), (b, y), width=5)
        bbox = d.textbbox((0, 0), label, font=diagram_font(21))
        mid = (a + b) / 2
        d.rectangle((mid-(bbox[2]-bbox[0])/2-7, y-36, mid+(bbox[2]-bbox[0])/2+7, y-7), fill="white")
        d.text((mid-(bbox[2]-bbox[0])/2, y-35), label, font=diagram_font(21), fill="#152238")
    d.text((700, 575), "หากไม่ใช่สมาชิก ระบบปฏิเสธการส่งข้อความ", font=diagram_font(20), fill="#B94A48")
    im.save(path, dpi=(180, 180)); return path


def save_activity_diagram(path_name, title, steps):
    path = DIAGRAM_DIR / path_name
    height = 260 + len(steps) * 155
    im = Image.new("RGB", (1200, height), "white")
    d = ImageDraw.Draw(im)
    centered_text(d, (100, 25, 1100, 95), title, diagram_font(32, True))
    y = 125
    rounded_box(d, (390, y, 810, y + 90), "เริ่มต้น", fill="#E8F1FF", font_size=25)
    prev = (600, y + 90)
    for index, step in enumerate(steps):
        y += 155
        arrow(d, prev, (600, y), width=5)
        fill = "#FFF4DF" if step.startswith("ตรวจสอบ") else "#F3F7FA"
        outline = "#9A6718" if step.startswith("ตรวจสอบ") else "#245780"
        rounded_box(d, (260, y, 940, y + 95), step, fill=fill, outline=outline, font_size=24)
        prev = (600, y + 95)
    y += 155
    arrow(d, prev, (600, y), width=5)
    rounded_box(d, (390, y, 810, y + 90), "สิ้นสุด", fill="#EAF7F0", outline="#2B7A58", font_size=25)
    im.save(path, dpi=(180, 180)); return path


def save_questionnaire_sequence_diagram():
    path = DIAGRAM_DIR / "questionnaire-sequence.png"
    im = Image.new("RGB", (1500, 1120), "white")
    d = ImageDraw.Draw(im)
    xs=[160,520,900,1270]; labels=["ผู้ใช้","Native iOS","Backend","ฐานข้อมูล"]
    for x,label in zip(xs,labels):
        rounded_box(d,(x-115,35,x+115,120),label,fill="#E8F1FF",font_size=23)
        d.line((x,120,x,1050),fill="#AAB4C0",width=3)
    steps=[(160,520,190,"เลือกความสนใจ งบ จำนวนสถานที่ และช่วงเวลา"),(520,900,330,"ส่งข้อมูล Preference พร้อม JWT"),(900,900,460,"ตรวจสอบและปรับรูปแบบข้อมูล"),(900,1270,610,"อัปเดต interests และ travelStyle"),(1270,900,750,"ยืนยันการบันทึก"),(900,520,880,"ส่งข้อมูลผู้ใช้ที่อัปเดตแล้ว"),(520,160,980,"แสดงสถานะบันทึกสำเร็จ")]
    for a,b,y,label in steps:
        if a == b:
            d.arc((a-95,y-25,a+95,y+55),start=270,end=90,fill="#245780",width=5); end=(a+95,y+15); arrow(d,end,(a,y+55),width=5)
        else: arrow(d,(a,y),(b,y),width=5)
        bbox=d.textbbox((0,0),label,font=diagram_font(21)); mid=(a+b)/2
        d.rectangle((mid-(bbox[2]-bbox[0])/2-6,y-35,mid+(bbox[2]-bbox[0])/2+6,y-7),fill="white")
        d.text((mid-(bbox[2]-bbox[0])/2,y-34),label,font=diagram_font(21),fill="#152238")
    im.save(path,dpi=(180,180)); return path


def save_admin_sequence_diagram():
    path = DIAGRAM_DIR / "admin-sequence.png"
    im = Image.new("RGB", (1600, 1330), "white")
    d = ImageDraw.Draw(im)
    xs=[150,500,850,1210,1480]; labels=["Admin","Admin Backoffice","Backend","ฐานข้อมูล","ผู้ใช้ที่เกี่ยวข้อง"]
    for x,label in zip(xs,labels):
        rounded_box(d,(x-110,35,x+110,120),label,fill="#E8F1FF",font_size=21)
        d.line((x,120,x,1250),fill="#AAB4C0",width=3)
    steps=[(150,500,180,"กรอกอีเมลและรหัสผ่าน"),(500,850,285,"POST Login"),(850,1210,390,"ตรวจบัญชีและ role"),(850,500,510,"ส่ง JWT และสิทธิ์ admin"),(500,850,640,"เรียก Dashboard พร้อม JWT"),(850,1210,750,"อ่าน User Trip และ Report"),(850,500,860,"ส่งข้อมูล Dashboard"),(150,500,980,"เลือกจัดการ User Trip หรือ Report"),(500,850,1090,"ส่งคำสั่งพร้อม JWT"),(850,1210,1180,"ตรวจ role และบันทึกการเปลี่ยนแปลง"),(850,1480,1260,"ส่งการแจ้งเตือนเมื่อกระบวนการกำหนดไว้")]
    for a,b,y,label in steps:
        arrow(d,(a,y),(b,y),width=5)
        bbox=d.textbbox((0,0),label,font=diagram_font(20)); mid=(a+b)/2
        d.rectangle((mid-(bbox[2]-bbox[0])/2-6,y-34,mid+(bbox[2]-bbox[0])/2+6,y-7),fill="white")
        d.text((mid-(bbox[2]-bbox[0])/2,y-33),label,font=diagram_font(20),fill="#152238")
    im.save(path,dpi=(180,180)); return path


def add_figure(doc, image_path, caption, width=6.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    r = c.add_run(caption)
    r.bold = True
    r.font.name = "Tahoma"
    r.font.size = Pt(13)


ARCHITECTURE_DIAGRAM = save_architecture_diagram()
USE_CASE_DIAGRAM = save_use_case_diagram()
OTP_SEQUENCE_DIAGRAM = save_otp_sequence_diagram()
ER_DIAGRAM = save_er_diagram()
MATCHING_FLOWCHART = save_matching_flowchart()
JOIN_SEQUENCE_DIAGRAM = save_join_sequence_diagram()
CREATE_TRIP_SEQUENCE_DIAGRAM = save_create_trip_sequence_diagram()
MESSAGE_SEQUENCE_DIAGRAM = save_message_sequence_diagram()
USER_ACTIVITY_DIAGRAM = save_activity_diagram("user-activity.png", "Activity Diagram ฝั่งผู้ใช้งาน", ["สมัครหรือเข้าสู่ระบบ", "ตอบแบบสอบถามและตั้งค่าโปรไฟล์", "โหลด Feed และดู Match Score", "ดูรายละเอียดหรือสร้างทริป", "ยืนยันเข้าร่วมทริป", "สนทนาและรับการแจ้งเตือน"])
SYSTEM_ACTIVITY_DIAGRAM = save_activity_diagram("system-activity.png", "Activity Diagram ฝั่งระบบ", ["ตรวจสอบ JWT และสิทธิ์", "อ่านข้อมูลผู้ใช้และทริป", "ตรวจสอบทริปที่สิ้นสุดหรือเต็ม", "คำนวณคะแนนย่อยจากข้อมูลที่มี", "คำนวณค่าเฉลี่ยและเรียงลำดับ", "ส่งข้อมูลกลับเป็น JSON"])
QUESTIONNAIRE_SEQUENCE_DIAGRAM = save_questionnaire_sequence_diagram()
ADMIN_SEQUENCE_DIAGRAM = save_admin_sequence_diagram()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph(style="Body Thai")
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.45)
    p.add_run(text)
    return p


def add_reference(doc, text):
    """Add a compact bibliography entry with a hanging indent."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Tahoma"
    r.font.size = Pt(13)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Body Thai")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.add_run(f"{index}. {item}")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_props.append(repeat_header)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, "1F4E78")
        if widths:
            cell.width = Inches(widths[i])
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            run.font.name = "Tahoma"
            run.font.size = Pt(14)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        set_cell_margins(cell)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if widths:
                cells[i].width = Inches(widths[i])
            if r_idx % 2:
                set_cell_shading(cells[i], "F3F7FA")
            for run in cells[i].paragraphs[0].runs:
                run.font.name = "Tahoma"
                run.font.size = Pt(13)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.2677)
section.page_height = Inches(11.6929)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1.35)
section.right_margin = Inches(0.85)
add_page_number(section.footer.paragraphs[0])

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Tahoma"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
normal.font.size = Pt(16)

if "Body Thai" not in styles:
    body_style = styles.add_style("Body Thai", WD_STYLE_TYPE.PARAGRAPH)
else:
    body_style = styles["Body Thai"]
body_style.font.name = "Tahoma"
body_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
body_style.font.size = Pt(16)
body_style.paragraph_format.line_spacing = 1.15
body_style.paragraph_format.space_after = Pt(6)
body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for level, size in ((1, 18), (2, 17), (3, 16)):
    style = styles[f"Heading {level}"]
    style.font.name = "Tahoma"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)

for name in ("List Bullet", "List Number"):
    styles[name].font.name = "Tahoma"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    styles[name].font.size = Pt(16)
    styles[name].paragraph_format.space_after = Pt(3)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("บทที่ 3")
r.bold = True
r.font.name = "Tahoma"
r.font.size = Pt(20)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run("วิธีการดำเนินการและการออกแบบระบบ")
r.bold = True
r.font.name = "Tahoma"
r.font.size = Pt(20)

add_body(doc, "บทนี้อธิบายกระบวนการวิเคราะห์ ออกแบบ พัฒนา และทดสอบแอปพลิเคชัน Go With Us ตามแนวทางวงจรการพัฒนาระบบ โดยขอบเขตของระบบประกอบด้วยแอปพลิเคชัน Native iOS สำหรับผู้ใช้งาน ระบบ Admin Backoffice สำหรับผู้ดูแล ระบบบริการส่วนหลัง และฐานข้อมูล PostgreSQL บน Supabase การแนะนำทริปใช้ข้อมูลความสนใจ งบประมาณ จำนวนสถานที่ต่อวัน และช่วงเวลาที่ชอบทำกิจกรรม เพื่อคำนวณคะแนนความเข้ากันได้ระหว่างผู้ใช้กับทริป")

add_heading(doc, "3.1 การวิเคราะห์ปัญหาและความต้องการของระบบ", 1)
add_body(doc, "ระบบแนะนำทริปทั่วไปมักแสดงรายการตามสถานที่ยอดนิยมโดยไม่ได้พิจารณาความต้องการเฉพาะบุคคล โครงงานนี้จึงออกแบบให้ผู้ใช้ระบุข้อมูลความชอบอย่างมีโครงสร้าง และนำข้อมูลดังกล่าวมาเปรียบเทียบกับคุณลักษณะของทริป ระบบใช้ Cosine Similarity เฉพาะข้อมูลที่อยู่ในรูปแบบเวกเตอร์หลายตัวเลือก ได้แก่ ความสนใจและช่วงเวลาที่ชอบ ส่วนงบประมาณและจำนวนสถานที่ต่อวันใช้สูตรคะแนนเชิงสัดส่วนและระยะห่าง")

add_heading(doc, "3.1.1 ความต้องการเชิงฟังก์ชัน", 2)
add_table(doc, ["กลุ่มระบบ", "ความสามารถที่กำหนด"], [
    ["บัญชีผู้ใช้", "สมัครสมาชิกด้วยชื่อ อีเมล และรหัสผ่าน ส่งรหัส OTP ทางอีเมล และสร้างบัญชีผู้ใช้จริงเมื่อยืนยัน OTP สำเร็จ"],
    ["การเข้าสู่ระบบ", "เข้าสู่ระบบด้วยอีเมลและรหัสผ่าน ตรวจสอบ JWT และสิทธิ์ user หรือ admin"],
    ["ข้อมูลส่วนตัว", "แสดงและแก้ไขชื่อ username รูปโปรไฟล์ แกลเลอรี เพศ ประวัติย่อ และความสนใจตามการตั้งค่าความเป็นส่วนตัว"],
    ["แบบสอบถาม", "บันทึกงบประมาณ จำนวนสถานที่ต่อวัน ช่วงเวลาที่ชอบ และหมวดหมู่ความสนใจ"],
    ["การค้นหาทริป", "เรียกข้อมูลทริปที่ยังไม่สิ้นสุด คำนวณ Match Score เรียงลำดับ และแสดงรายละเอียดคะแนนย่อย"],
    ["การสร้างทริป", "บันทึกชื่อทริป จังหวัด วันเดินทาง งบประมาณ จำนวนผู้ร่วมเดินทาง หมวดหมู่ รูปภาพ และแผนกิจกรรม"],
    ["สมาชิกทริป", "บันทึกสถานะสนใจหรือยืนยันเข้าร่วม และป้องกันการเข้าร่วมซ้ำด้วยคู่ tripId และ userId"],
    ["การสนทนา", "ส่งข้อความหรือรูปภาพ โดยรองรับข้อความกลุ่มผ่าน tripId และข้อความส่วนตัวผ่าน receiverId"],
    ["การแจ้งเตือน", "แสดงประกาศหรือการแจ้งเตือนเฉพาะผู้ใช้และทั้งระบบ พร้อมสถานะอ่านแล้ว"],
    ["รายงานผู้ใช้", "ส่งรายงาน ระบุเหตุผล และติดตามสถานะ pending reviewed หรือ resolved"],
    ["Admin Backoffice", "ตรวจสอบสิทธิ์ admin และจัดการข้อมูลผู้ใช้ ทริป รายงาน การยืนยันตัวตน และประกาศระบบ"],
], [1.35, 4.7])

add_heading(doc, "3.1.2 ความต้องการเชิงไม่ใช่ฟังก์ชัน", 2)
add_table(doc, ["ด้าน", "ข้อกำหนด"], [
    ["ความปลอดภัย", "แฮชรหัสผ่านก่อนบันทึก ใช้ JWT ควบคุมสิทธิ์ ใช้ HTTPS และขอสิทธิ์กล้องเฉพาะขั้นตอนที่จำเป็น"],
    ["ความถูกต้อง", "ตรวจสอบข้อมูลนำเข้า ป้องกันผู้เข้าร่วมซ้ำ ไม่แสดงทริปที่สิ้นสุดหรือเต็ม และจำกัดคะแนนให้อยู่ระหว่าง 0–100"],
    ["การใช้งาน", "ออกแบบแอป iOS ตามแนวทางของ Apple รองรับภาษาไทยและอังกฤษ และแสดงสถานะโหลดหรือข้อผิดพลาดอย่างชัดเจน"],
    ["การบำรุงรักษา", "แยกส่วน iOS, Admin Backoffice, Backend และฐานข้อมูล พร้อมรวมตรรกะคำนวณไว้ที่ Backend"],
    ["ประสิทธิภาพ", "บันทึกสภาพแวดล้อมและวัดเวลาตอบสนองจริงก่อนสรุปผล ไม่กำหนดตัวเลขผ่านเกณฑ์โดยไม่มีผลทดสอบรองรับ"],
], [1.35, 4.7])

add_heading(doc, "3.2 เทคโนโลยีและสถาปัตยกรรมระบบ", 1)
add_heading(doc, "3.2.1 เทคโนโลยีที่ใช้", 2)
add_table(doc, ["ส่วนระบบ", "เทคโนโลยี", "หน้าที่"], [
    ["Native iOS", "Swift, SwiftUI, MapKit, Vision", "ส่วนติดต่อผู้ใช้ แผนที่ กล้อง และการตรวจจับใบหน้าแบบทำงานบนอุปกรณ์"],
    ["Admin Backoffice", "TypeScript, React, Vite, Tailwind CSS", "ส่วนติดต่อสำหรับผู้ดูแลระบบผ่านเว็บเบราว์เซอร์"],
    ["Backend", "Node.js, Express, Prisma", "API ตรรกะธุรกิจ การตรวจสอบสิทธิ์ และ Smart Matching"],
    ["ฐานข้อมูล", "PostgreSQL บน Supabase", "จัดเก็บบัญชีผู้ใช้ ทริป ผู้เข้าร่วม ข้อความ การแจ้งเตือน และรายงาน"],
    ["บริการภายนอก", "Gemini API และบริการอีเมล", "สร้างร่างแผนการเดินทางและส่ง OTP"],
], [1.3, 2.15, 2.6])

add_heading(doc, "3.2.2 สถาปัตยกรรมแบบ Client Server", 2)
add_body(doc, "แอปพลิเคชัน iOS และ Admin Backoffice ทำหน้าที่เป็น Client และเรียกใช้ Backend ผ่าน REST API ในรูปแบบ JSON ตามแนวคิดที่แยกผู้ร้องขอบริการออกจากผู้ให้บริการ (van Leeuwen, 1988) Backend ตรวจสอบ JWT และสิทธิ์ของผู้ใช้ ประมวลผลกฎทางธุรกิจ เรียก Prisma เพื่ออ่านหรือเขียนข้อมูลใน PostgreSQL และส่งผลลัพธ์กลับไปยัง Client")
add_figure(doc, ARCHITECTURE_DIAGRAM, "ภาพที่ 3.1 สถาปัตยกรรมระบบ Go With Us")

add_heading(doc, "3.3 ขั้นตอนการดำเนินงาน", 1)
add_numbered(doc, [
    "วิเคราะห์ปัญหา ขอบเขต ผู้ใช้งาน และข้อมูลที่จำเป็นต่อการแนะนำทริป",
    "ออกแบบหน้าจอ กระบวนการใช้งาน API โครงสร้างฐานข้อมูล และสูตร Smart Matching",
    "พัฒนา Native iOS, Admin Backoffice และ Backend แยกตามหน้าที่",
    "เชื่อมต่อฐานข้อมูลและบริการภายนอก พร้อมตรวจสอบสิทธิ์ด้วย JWT",
    "ทดสอบระดับหน่วย การเชื่อมต่อ ระบบรวม และการยอมรับของผู้ใช้",
    "บันทึกผลทดสอบจริง วิเคราะห์ข้อจำกัด และปรับปรุงระบบ",
])

add_heading(doc, "3.4 การวิเคราะห์และออกแบบระบบ", 1)
add_heading(doc, "3.4.1 ผู้ใช้งานระบบและกรณีใช้งาน", 2)
add_table(doc, ["Actor", "กรณีใช้งานหลัก"], [
    ["ผู้เยี่ยมชม", "สมัครสมาชิก ยืนยัน OTP และเข้าสู่ระบบ"],
    ["ผู้ใช้งาน", "ตอบแบบสอบถาม แก้ไขโปรไฟล์ ค้นหาและสร้างทริป เข้าร่วมทริป สนทนา และรายงานผู้ใช้"],
    ["ผู้ดูแลระบบ", "เข้าสู่ Admin Backoffice ตรวจสอบ Dashboard ผู้ใช้ ทริป รายงาน การยืนยันตัวตน และประกาศ"],
    ["บริการภายนอก", "ส่ง OTP สร้างร่างแผนเดินทาง และให้ข้อมูลแผนที่"],
], [1.35, 4.7])
add_figure(doc, USE_CASE_DIAGRAM, "ภาพที่ 3.2 แผนภาพกรณีใช้งานของผู้ใช้และผู้ดูแลระบบ", width=5.7)

add_heading(doc, "3.4.2 Activity Diagram ฝั่งผู้ใช้งาน", 2)
add_body(doc, "แผนภาพฝั่งผู้ใช้งานแสดงเส้นทางหลักตั้งแต่เข้าสู่ระบบ ระบุความต้องการ โหลดรายการทริป ดูรายละเอียด สร้างหรือเข้าร่วมทริป ไปจนถึงการสนทนาและรับการแจ้งเตือน")
add_figure(doc, USER_ACTIVITY_DIAGRAM, "ภาพที่ 3.3 Activity Diagram ฝั่งผู้ใช้งาน", width=4.9)

add_heading(doc, "3.4.3 Activity Diagram ฝั่งระบบ", 2)
add_body(doc, "แผนภาพฝั่งระบบแสดงการตรวจสิทธิ์ การอ่านข้อมูล การคัดกรองทริป การคำนวณคะแนน และการส่งผลลัพธ์กลับแก่แอปพลิเคชัน")
add_figure(doc, SYSTEM_ACTIVITY_DIAGRAM, "ภาพที่ 3.4 Activity Diagram ฝั่งระบบ", width=4.9)

add_heading(doc, "3.4.4 ลำดับการสมัครสมาชิกและยืนยัน OTP", 2)
add_numbered(doc, [
    "ผู้ใช้กรอกชื่อ อีเมล และรหัสผ่าน",
    "Backend ตรวจสอบข้อมูล แฮชรหัสผ่าน และบันทึกไว้ใน PendingRegistration",
    "ระบบส่ง OTP ไปยังอีเมลและกำหนดเวลาหมดอายุ",
    "ผู้ใช้กรอก OTP หรือย้อนกลับเพื่อแก้ไขข้อมูล",
    "เมื่อ OTP ถูกต้อง Backend จึงสร้างข้อมูลใน User และลบ PendingRegistration",
    "หาก OTP ไม่ถูกต้องหรือหมดอายุ ระบบไม่สร้างบัญชีผู้ใช้จริง",
])
add_figure(doc, OTP_SEQUENCE_DIAGRAM, "ภาพที่ 3.5 แผนภาพลำดับการสมัครสมาชิกและยืนยัน OTP", width=5.9)

add_heading(doc, "3.4.5 ลำดับการประมวลผลหน้าแนะนำทริป", 2)
add_numbered(doc, [
    "แอปพลิเคชันส่ง JWT เพื่อเรียกรายการทริปที่ตรงกับผู้ใช้",
    "Backend อ่าน interests และ travelStyle ของผู้ใช้",
    "Backend อ่านทริปที่ยังไม่สิ้นสุด พร้อมข้อมูลผู้สร้างและผู้เข้าร่วม",
    "ระบบตัดทริปที่เต็มหรือสิ้นสุดออกจากการแนะนำ",
    "ระบบคำนวณคะแนนย่อยทั้งสี่ปัจจัยจากข้อมูลที่มี",
    "ระบบหาค่าเฉลี่ยของคะแนนย่อย เรียงจากมากไปน้อย และส่งผลเป็น JSON",
    "แอปพลิเคชันแสดง Match Score และรายละเอียดแต่ละปัจจัยบนหน้าจอ",
])
add_heading(doc, "3.4.6 ลำดับการตอบแบบสอบถามและบันทึก Preference", 2)
add_body(doc, "ผู้ใช้เลือกความสนใจ งบประมาณสูงสุด จำนวนสถานที่ต่อวัน และช่วงเวลาที่ชอบ แอปส่งข้อมูลพร้อม JWT ให้ Backend ตรวจสอบก่อนอัปเดต interests และ travelStyle ใน users")
add_figure(doc, QUESTIONNAIRE_SEQUENCE_DIAGRAM, "ภาพที่ 3.6 แผนภาพลำดับการตอบแบบสอบถามและบันทึก Preference", width=6.0)

add_heading(doc, "3.4.7 ลำดับการยืนยันเข้าร่วมทริป", 2)
add_figure(doc, JOIN_SEQUENCE_DIAGRAM, "ภาพที่ 3.7 แผนภาพลำดับการยืนยันเข้าร่วมทริป", width=6.0)

add_heading(doc, "3.4.8 ลำดับการสร้างทริป", 2)
add_body(doc, "ผู้ใช้ต้องผ่านการตรวจสอบ JWT และส่งชื่อทริป จุดหมาย และวันเริ่มต้น Backend ตรวจสอบผู้สร้าง คำนวณจำนวนสถานที่ต่อวันจาก itinerary และสร้างทริปพร้อมบันทึกผู้สร้างเป็น Participant โดยอัตโนมัติ การสร้าง embedding ผ่าน Gemini เป็นกระบวนการเสริมและไม่ขัดขวางการสร้างทริปเมื่อบริการภายนอกไม่พร้อมใช้งาน")
add_figure(doc, CREATE_TRIP_SEQUENCE_DIAGRAM, "ภาพที่ 3.8 แผนภาพลำดับการสร้างทริป", width=6.0)

add_heading(doc, "3.4.9 ลำดับการส่งข้อความในทริป", 2)
add_body(doc, "ผู้ส่งต้องเป็น Participant ของทริปและต้องส่งข้อความหรือรูปภาพอย่างน้อยหนึ่งรายการ Backend บันทึก Message แล้วส่งข้อมูลกลับแก่ผู้ส่ง พร้อมแจ้งสมาชิกคนอื่นผ่าน WebSocket การแจ้งเตือนในฐานข้อมูล และ Push Notification เมื่อมี FCM Token")
add_figure(doc, MESSAGE_SEQUENCE_DIAGRAM, "ภาพที่ 3.9 แผนภาพลำดับการส่งข้อความในทริป", width=6.0)

add_heading(doc, "3.4.10 ลำดับการเข้าสู่ระบบและจัดการข้อมูลของ Admin", 2)
add_body(doc, "Admin เข้าสู่ระบบด้วยบัญชีที่มี role เป็น admin Backend ออก JWT หลังตรวจสอบข้อมูลสำเร็จ จากนั้น Admin Backoffice จึงโหลด Dashboard และส่งคำสั่งจัดการ User, Trip หรือ Report โดย Backend ตรวจ role ทุกครั้งก่อนแก้ไขฐานข้อมูล")
add_figure(doc, ADMIN_SEQUENCE_DIAGRAM, "ภาพที่ 3.10 แผนภาพลำดับการเข้าสู่ระบบและจัดการข้อมูลของ Admin", width=6.0)

add_heading(doc, "3.4.11 การออกแบบฐานข้อมูล", 2)
add_body(doc, "ฐานข้อมูลใช้โครงสร้างเชิงสัมพันธ์ร่วมกับข้อมูล JSON และอาร์เรย์ในฟิลด์ที่เหมาะสม ข้อมูลแบบสอบถามไม่ได้แยกเป็นตารางใหม่ แต่เก็บความสนใจไว้ใน User.interests และเก็บงบประมาณ จำนวนสถานที่ต่อวัน และช่วงเวลาที่ชอบไว้ใน User.travelStyle")
add_figure(doc, ER_DIAGRAM, "ภาพที่ 3.11 แผนภาพความสัมพันธ์ของข้อมูลหลัก", width=6.0)
add_table(doc, ["ตาราง", "วัตถุประสงค์", "ความสัมพันธ์สำคัญ"], [
    ["users", "บัญชี โปรไฟล์ สิทธิ์ ความสนใจ travelStyle และสถานะยืนยัน", "ผู้ใช้หนึ่งคนสร้างทริป ส่งข้อความ เข้าร่วมทริป และส่งหรือถูกรายงานได้หลายรายการ"],
    ["pending_registrations", "ข้อมูลสมัครชั่วคราวและ OTP ก่อนสร้างบัญชีจริง", "เชื่อมกระบวนการด้วยอีเมล และถูกลบเมื่อยืนยันสำเร็จ"],
    ["trips", "ข้อมูลทริป งบ หมวดหมู่ เวลา รูป และ itinerary", "แต่ละทริปมี creator หนึ่งคน และมี participants กับ messages หลายรายการ"],
    ["participants", "ตารางกลางระหว่าง users กับ trips", "กำหนดคู่ tripId และ userId ไม่ให้ซ้ำ"],
    ["messages", "ข้อความกลุ่มหรือส่วนตัวและรูปภาพ", "เชื่อม sender เสมอ และเชื่อม receiver หรือ trip ตามชนิดข้อความ"],
    ["notifications", "ประกาศ แจ้งเตือนทริป หรือแจ้งเตือนระบบ", "userId ว่างหมายถึงประกาศถึงทุกคน"],
    ["user_reports", "รายงานผู้ใช้และสถานะดำเนินการ", "เชื่อม reporter และ reported ไปยัง users"],
], [1.45, 2.45, 2.15])

add_heading(doc, "3.4.12 Data Dictionary ของข้อมูลทั้งหมด", 2)
add_body(doc, "ตารางต่อไปนี้แจกแจงทุกคอลัมน์ที่บันทึกจริงตาม Prisma Schema โดยไม่รวม relation field ที่ Prisma ใช้เชื่อมออบเจ็กต์แต่ไม่ได้สร้างเป็นคอลัมน์เพิ่มเติมในฐานข้อมูล", first_line=False)
add_table(doc, ["ตาราง", "ฟิลด์", "ชนิดข้อมูล", "Key หรือ Constraint", "รายละเอียด"], [
    ["users", "id", "String UUID", "PK", "รหัสผู้ใช้"],
    ["users", "username", "String?", "Unique", "ชื่อบัญชีแบบ @"],
    ["users", "usernameUpdatedAt", "DateTime?", "-", "เวลาที่เปลี่ยน username ล่าสุด"],
    ["users", "name", "String", "Required", "ชื่อที่แสดง"],
    ["users", "email", "String", "Unique", "อีเมลสำหรับเข้าสู่ระบบ"],
    ["users", "password", "String", "Required", "รหัสผ่านที่ผ่านการแฮช"],
    ["users", "role", "String", "Default user", "สิทธิ์ user หรือ admin"],
    ["users", "isBanned", "Boolean", "Default false", "สถานะระงับบัญชี"],
    ["users", "gender", "String?", "-", "เพศ"],
    ["users", "age", "Int?", "-", "อายุ"],
    ["users", "bio", "Text?", "-", "คำแนะนำตัว"],
    ["users", "birthDate", "DateTime?", "-", "วันเกิด"],
    ["users", "profileImage", "Text?", "-", "URL หรือ Base64 รูปหลัก"],
    ["users", "gallery", "String[]", "Default []", "รูปเพิ่มเติม"],
    ["users", "isProfilePublic", "Boolean", "Default true", "เปิดเผยโปรไฟล์"],
    ["users", "showGender", "Boolean", "Default true", "แสดงเพศ"],
    ["users", "showAge", "Boolean", "Default true", "แสดงอายุ"],
    ["users", "showBio", "Boolean", "Default true", "แสดงคำแนะนำตัว"],
    ["users", "showInterests", "Boolean", "Default true", "แสดงความสนใจ"],
    ["users", "showEmail", "Boolean", "Default false", "แสดงอีเมล"],
    ["users", "createdAt", "DateTime", "Default now", "เวลาสร้างบัญชี"],
    ["users", "updatedAt", "DateTime", "Auto update", "เวลาแก้ไขล่าสุด"],
    ["users", "interests", "String[]", "Default []", "หมวดความสนใจ"],
    ["users", "travelStyle", "JSON?", "-", "งบ จำนวนสถานที่ และช่วงเวลา"],
    ["users", "embedding", "JSON?", "-", "เวกเตอร์เสริมสำหรับ AI"],
    ["users", "fcmToken", "String?", "-", "โทเคน Push Notification"],
    ["users", "isEmailVerified", "Boolean", "Default false", "สถานะยืนยันอีเมล"],
    ["users", "otpCode", "String?", "-", "OTP เดิมที่ยังอยู่ใน schema"],
    ["users", "otpExpiresAt", "DateTime?", "-", "เวลาหมดอายุ OTP เดิม"],
    ["users", "isVerified", "Boolean", "Default false", "สถานะยืนยันตัวตน"],
    ["users", "verificationStatus", "String", "Default unverified", "สถานะตรวจสอบตัวตน"],
    ["users", "idCardImage", "Text?", "-", "ภาพบัตรเดิมที่ยังอยู่ใน schema"],
    ["users", "faceScanImage", "Text?", "-", "ภาพ Selfie หลังตรวจ liveness"],
    ["pending_registrations", "email", "String", "PK", "อีเมลที่รอยืนยัน"],
    ["pending_registrations", "name", "String", "Required", "ชื่อผู้สมัคร"],
    ["pending_registrations", "password_hash", "String", "Required", "รหัสผ่านที่แฮชแล้ว"],
    ["pending_registrations", "otp_code", "String", "Required", "รหัส OTP"],
    ["pending_registrations", "otp_expires_at", "DateTime", "Required", "เวลาหมดอายุ OTP"],
    ["pending_registrations", "created_at", "DateTime", "Default now", "เวลาสร้างรายการ"],
    ["pending_registrations", "updated_at", "DateTime", "Auto update", "เวลาแก้ไขล่าสุด"],
    ["trips", "id", "String UUID", "PK", "รหัสทริป"],
    ["trips", "title", "String", "Required", "ชื่อทริป"],
    ["trips", "destination", "String", "Index", "จุดหมายหรือจังหวัด"],
    ["trips", "description", "String?", "-", "รายละเอียดทริป"],
    ["trips", "start_date", "DateTime", "Required", "วันเริ่มต้น"],
    ["trips", "end_date", "DateTime?", "-", "วันสิ้นสุด"],
    ["trips", "budget", "Int", "Default 1000", "งบประมาณทริป"],
    ["trips", "budget_type", "String", "Default per_person", "ประเภทงบประมาณ"],
    ["trips", "max_participants", "Int", "Default 10", "จำนวนสมาชิกสูงสุด"],
    ["trips", "activity_style", "Int?", "Default 5", "จำนวนสถานที่เฉลี่ยต่อวัน"],
    ["trips", "time_of_day", "String[]", "Default []", "ช่วงเวลาของกิจกรรม"],
    ["trips", "budget_rating", "Int?", "Default 5", "ระดับงบเดิมที่ยังอยู่ใน schema"],
    ["trips", "category", "String?", "Index", "หมวดหลัก"],
    ["trips", "interest_tags", "String[]", "Default []", "หมวดเสริม"],
    ["trips", "is_public", "Boolean", "Default true", "การมองเห็นทริป"],
    ["trips", "imageUrl", "Text?", "-", "รูปหน้าปก"],
    ["trips", "gallery", "String[]", "Default []", "รูปเพิ่มเติม"],
    ["trips", "itinerary", "JSON?", "-", "แผนรายวันและสถานที่"],
    ["trips", "embedding", "JSON?", "-", "เวกเตอร์เสริม"],
    ["trips", "creator_id", "String", "FK users.id", "ผู้สร้างทริป"],
    ["trips", "created_at", "DateTime", "Default now", "เวลาสร้างทริป"],
    ["trips", "updated_at", "DateTime", "Auto update", "เวลาแก้ไขล่าสุด"],
    ["trips", "summary", "Text?", "-", "สรุปทริป"],
    ["trips", "groupAnalysis", "Text?", "-", "ผลวิเคราะห์กลุ่ม"],
    ["participants", "id", "String UUID", "PK", "รหัสรายการเข้าร่วม"],
    ["participants", "trip_id", "String", "FK trips.id", "ทริปที่เข้าร่วม"],
    ["participants", "user_id", "String", "FK users.id", "ผู้เข้าร่วม"],
    ["participants", "name", "String", "Required", "ชื่อที่บันทึกในทริป"],
    ["participants", "interests", "String[]", "Required", "ความสนใจของสมาชิก"],
    ["participants", "status", "String", "Default going", "interested หรือ going"],
    ["participants", "joined_at", "DateTime", "Default now", "เวลาเข้าร่วม"],
    ["messages", "id", "String UUID", "PK", "รหัสข้อความ"],
    ["messages", "content", "Text", "Required", "ข้อความ"],
    ["messages", "sender_id", "String", "FK users.id", "ผู้ส่ง"],
    ["messages", "receiver_id", "String?", "FK users.id", "ผู้รับข้อความส่วนตัว"],
    ["messages", "trip_id", "String?", "FK trips.id", "ห้องข้อความกลุ่ม"],
    ["messages", "created_at", "DateTime", "Default now", "เวลาส่ง"],
    ["messages", "is_read", "Boolean", "Default false", "สถานะอ่าน"],
    ["messages", "imageUrl", "Text?", "-", "รูปภาพในข้อความ"],
    ["notifications", "id", "String UUID", "PK", "รหัสการแจ้งเตือน"],
    ["notifications", "title", "String", "Required", "หัวข้อ"],
    ["notifications", "message", "Text?", "-", "รายละเอียด"],
    ["notifications", "type", "String", "Default alert", "ประเภทการแจ้งเตือน"],
    ["notifications", "target_id", "String?", "Index", "รหัสเป้าหมาย เช่น ทริป"],
    ["notifications", "user_id", "String?", "Index", "ผู้รับ หรือว่างเมื่อส่งทั้งหมด"],
    ["notifications", "is_read", "Boolean", "Default false", "สถานะอ่าน"],
    ["notifications", "created_at", "DateTime", "Default now", "เวลาสร้าง"],
    ["user_reports", "id", "String UUID", "PK", "รหัสรายงาน"],
    ["user_reports", "reason", "Text", "Required", "เหตุผลการรายงาน"],
    ["user_reports", "status", "String", "Default pending", "สถานะดำเนินการ"],
    ["user_reports", "createdAt", "DateTime", "Default now", "เวลารายงาน"],
    ["user_reports", "reporterId", "String", "FK users.id", "ผู้รายงาน"],
    ["user_reports", "reportedId", "String", "FK users.id", "ผู้ถูกรายงาน"],
    ["user_matches", "id", "String UUID", "PK", "รหัสข้อมูลเดิม"],
    ["user_matches", "liker_id", "String", "Unique pair", "ผู้กดเลือกในระบบเดิม"],
    ["user_matches", "liked_id", "String", "Unique pair", "ผู้ถูกเลือกในระบบเดิม"],
    ["user_matches", "status", "String", "Default like", "สถานะเดิม"],
    ["user_matches", "is_mutual", "Boolean", "Default false", "สถานะเลือกตรงกันเดิม"],
    ["user_matches", "created_at", "DateTime", "Default now", "เวลาสร้างข้อมูลเดิม"],
], [0.95, 1.25, 1.05, 1.35, 2.0])
add_body(doc, "หมายเหตุ ตาราง user_matches และฟิลด์ที่ระบุว่าเป็นข้อมูลเดิมยังปรากฏใน Prisma Schema แต่ไม่ได้ใช้เป็นส่วนหนึ่งของ Smart Matching ระหว่างผู้ใช้กับทริปในเวอร์ชันปัจจุบัน", first_line=False)

add_heading(doc, "3.5 การออกแบบระบบ Smart Matching", 1)
add_heading(doc, "3.5.1 ข้อมูลนำเข้า", 2)
add_table(doc, ["ปัจจัย", "ข้อมูลผู้ใช้", "ข้อมูลทริป"], [
    ["ความสนใจ", "เลือกได้หลายหมวดหมู่", "หมวดหลักและหมวดเสริมไม่เกิน 3 หมวด"],
    ["งบประมาณ", "งบประมาณสูงสุดต่อคน", "งบประมาณทริปต่อคน"],
    ["จำนวนสถานที่ต่อวัน", "เลือกจำนวนที่ต้องการ 1–10", "ค่าเฉลี่ยจากจำนวนกิจกรรมใน itinerary ต่อจำนวนวัน"],
    ["ช่วงเวลา", "เช้า กลางวัน เย็น หรือกลางคืน เลือกได้หลายช่วง", "ช่วงเวลาที่กำหนดในทริป"],
], [1.45, 2.3, 2.3])

add_heading(doc, "3.5.2 คะแนนความสนใจ", 2)
add_body(doc, "ระบบกำหนดลำดับหมวดหมู่ร่วมกัน 13 หมวด และแปลงรายการที่เลือกเป็นเวกเตอร์แบบ Multi-hot โดยเลือกเท่ากับ 1 และไม่เลือกเท่ากับ 0 จากนั้นปรับเวกเตอร์ให้มีความยาวหนึ่งหน่วยและคำนวณ Cosine Similarity ตามแนวคิด Vector Space Model (Salton, Wong, & Yang, 1975) ผลลัพธ์คูณ 100 เพื่อแสดงเป็นคะแนนความสนใจ")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("InterestScore = CosineSimilarity(UserInterestVector, TripInterestVector) × 100").bold = True

add_heading(doc, "3.5.3 คะแนนช่วงเวลา", 2)
add_body(doc, "ระบบแปลงช่วงเวลา morning, noon, evening และ night เป็นเวกเตอร์ Multi-hot และคำนวณด้วย Cosine Similarity เช่นเดียวกับความสนใจ การปรับความยาวเวกเตอร์ช่วยลดความได้เปรียบจากการเลือกตัวเลือกจำนวนมาก")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("TimeScore = CosineSimilarity(UserTimeVector, TripTimeVector) × 100").bold = True

add_heading(doc, "3.5.4 คะแนนงบประมาณ", 2)
add_body(doc, "หากงบของผู้ใช้มากกว่าหรือเท่ากับงบทริป ระบบให้ 100 คะแนน หากงบผู้ใช้น้อยกว่า ระบบให้คะแนนตามสัดส่วนจริง วิธีนี้ทำให้ผู้ใช้งบ 400 บาทกับทริป 1,000 บาทได้ 40 คะแนน และผู้ใช้ที่มีงบ 0 บาทกับทริป 1,000 บาทได้ 0 คะแนน")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("BudgetScore = 100 เมื่อ UserBudget ≥ TripBudget มิฉะนั้น (UserBudget ÷ TripBudget) × 100").bold = True

add_heading(doc, "3.5.5 คะแนนจำนวนสถานที่ต่อวัน", 2)
add_body(doc, "ผู้ใช้และทริปใช้ช่วงค่าเดียวกันตั้งแต่ 1–10 สถานที่ต่อวัน ระบบหาค่าความต่างสัมบูรณ์และเทียบกับระยะห่างสูงสุดซึ่งเท่ากับ 9 หากจำนวนเท่ากันได้ 100 คะแนน หากต่างกัน 1 สถานที่จะได้ประมาณ 89 คะแนน และหากต่างกันมากที่สุดได้ 0 คะแนน")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("ActivityScore = max(0, 1 − |UserActivity − TripActivity| ÷ 9) × 100").bold = True

add_heading(doc, "3.5.6 คะแนนรวม", 2)
add_body(doc, "คะแนนรวมคำนวณจากค่าเฉลี่ยเลขคณิตของคะแนนย่อยที่มีข้อมูลเพียงพอสำหรับการคำนวณในทริปรายการนั้น และปัดผลลัพธ์เป็นจำนวนเต็ม")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("MatchScore = ผลรวมคะแนนย่อยที่คำนวณได้ ÷ จำนวนปัจจัยที่มีข้อมูล").bold = True
add_figure(doc, MATCHING_FLOWCHART, "ภาพที่ 3.12 Flowchart การคำนวณ Smart Matching", width=5.7)

add_heading(doc, "3.5.7 ตัวอย่างการคำนวณ", 2)
add_table(doc, ["ปัจจัย", "คะแนนตัวอย่าง", "เหตุผล"], [
    ["ความสนใจ", "82", "ผู้ใช้และทริปมีหมวดหมู่ร่วมกันหลายรายการ"],
    ["งบประมาณ", "40", "งบผู้ใช้ 400 บาท เทียบกับงบทริป 1,000 บาท"],
    ["จำนวนสถานที่ต่อวัน", "89", "ผู้ใช้ต้องการ 2 สถานที่ แต่ทริปเฉลี่ย 3 สถานที่ต่อวัน"],
    ["ช่วงเวลา", "71", "เวกเตอร์ช่วงเวลามีบางช่วงที่เลือกตรงกัน"],
    ["คะแนนรวม", "71", "(82 + 40 + 89 + 71) ÷ 4 = 70.5 ปัดเป็น 71"],
], [1.7, 1.3, 3.05])

add_heading(doc, "3.5.8 การแปลผลคะแนน", 2)
add_body(doc, "ช่วงคะแนนต่อไปนี้เป็นเกณฑ์การนำเสนอที่ผู้พัฒนากำหนดเพื่อช่วยให้ผู้ใช้เข้าใจผลลัพธ์ ไม่ใช่ข้อพิสูจน์ว่าผู้ใช้จะเดินทางร่วมกันโดยไม่มีความขัดแย้ง")
add_table(doc, ["คะแนน", "ข้อความแสดงผล"], [
    ["76–100", "มีคุณลักษณะสอดคล้องกันมาก"],
    ["51–75", "มีคุณลักษณะสอดคล้องกันระดับปานกลาง"],
    ["26–50", "มีคุณลักษณะสอดคล้องกันบางส่วน"],
    ["0–25", "มีคุณลักษณะสอดคล้องกันน้อย"],
], [1.5, 4.55])

add_heading(doc, "3.6 การออกแบบการยืนยันตัวตน", 1)
add_body(doc, "การยืนยันอีเมลเกิดขึ้นระหว่างสมัครสมาชิก โดยระบบยังไม่สร้าง User จนกว่า OTP จะถูกต้อง ส่วนการตรวจจับใบหน้าใช้กล้องหน้าและ Apple Vision สำหรับตรวจตำแหน่งใบหน้าและการทำ Active Liveness ตามคำสั่งแบบสุ่ม เช่น หันซ้าย หันขวา หรือยิ้ม ระบบส่วนนี้ไม่ใช่ Face ID ของ Apple และไม่ควรอธิบายว่าเป็นการพิสูจน์ว่าใบหน้าตรงกับบุคคลในเอกสารราชการ เว้นแต่มีระบบเปรียบเทียบอัตลักษณ์เพิ่มเติมจริง")

add_heading(doc, "3.7 แผนการทดสอบและประเมินผล", 1)
add_table(doc, ["ระดับ", "สิ่งที่ทดสอบ", "หลักฐานที่ต้องบันทึก"], [
    ["Unit Testing", "สูตรคะแนน ฟังก์ชันตรวจข้อมูล และกรณีขอบเขต", "Input Expected Result Actual Result และสถานะ"],
    ["Integration Testing", "iOS กับ API, Admin กับ API, Prisma กับ PostgreSQL และบริการภายนอก", "Endpoint สถานะ HTTP และข้อมูลตอบกลับ"],
    ["System Testing", "กระบวนการสมัคร สร้างทริป เข้าร่วม แชต Matching และ Admin", "Test Case ภาพหน้าจอ และข้อผิดพลาด"],
    ["UAT", "ภารกิจที่ผู้ใช้จริงต้องทำ", "จำนวนผู้ทดสอบ อัตราสำเร็จ เวลา และปัญหา"],
    ["SUS", "แบบประเมินการใช้งาน 10 ข้อตาม Brooke (1996)", "คำตอบรายคน วิธีคำนวณ และคะแนนเฉลี่ย"],
    ["Performance", "เวลาโหลด Feed เวลา API และการคำนวณ Matching", "อุปกรณ์ เครือข่าย จำนวนครั้ง ค่าเฉลี่ย ต่ำสุด และสูงสุด"],
], [1.3, 2.75, 2.0])

add_heading(doc, "3.8 สรุปวิธีการดำเนินงาน", 1)
add_body(doc, "ระบบ Go With Us แยกแอปพลิเคชัน iOS, Admin Backoffice, Backend และฐานข้อมูลอย่างชัดเจน ข้อมูล interests และ travelStyle ใช้คำนวณความเข้ากันได้ระหว่างผู้ใช้กับทริปตามชนิดของแต่ละปัจจัย ส่วนความถูกต้องและประสิทธิภาพต้องสรุปจากผลทดสอบจริงในบทที่ 4")

doc.add_page_break()
add_heading(doc, "บรรณานุกรม", 1)
add_reference(doc, "Brooke, J. (1996). SUS: A quick and dirty usability scale. In P. W. Jordan, B. Thomas, B. A. Weerdmeester, and I. L. McClelland (Eds.), Usability evaluation in industry (pp. 189–194). Taylor and Francis.")
add_reference(doc, "Salton, G., Wong, A., and Yang, C. S. (1975). A vector space model for automatic indexing. Communications of the ACM, 18(11), 613–620. https://doi.org/10.1145/361219.361220")
add_reference(doc, "van Leeuwen, J. (1988). The client server model in distributed computing (RUU-CS-88-09). Utrecht University.")

doc.core_properties.title = "บทที่ 3 วิธีการดำเนินการและการออกแบบระบบ"
doc.core_properties.subject = "ฉบับแก้ไขทดลองสำหรับโครงการ Go With Us"
doc.core_properties.author = "Go With Us"

# Ensure LibreOffice and Word both use a Thai-capable complex-script font.
def normalize_run_font(run):
    run.font.name = "Arial Unicode MS"
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn("w:" + attr), "Arial Unicode MS")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "th-TH")
    lang.set(qn("w:eastAsia"), "th-TH")
    lang.set(qn("w:bidi"), "th-TH")


for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        normalize_run_font(run)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    normalize_run_font(run)

doc.save(OUT)
print(OUT)
