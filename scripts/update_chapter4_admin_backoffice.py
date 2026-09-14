from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path("/Users/worakanp/Desktop/รายงานบทที่4_วางข้อความใหม่.docx")
OUTPUT = Path("/Users/worakanp/Desktop/Go-with-us-1/docs/รายงานบทที่4_เพิ่มรูป_Admin_Backoffice.docx")
IMAGE_DIR = Path("/Users/worakanp/Desktop/Go-with-us-1/reports/backofficeรูปหลังบ้าน")


SECTIONS = [
    {
        "paragraph": 629,
        "image_paragraph": 630,
        "caption_paragraph": 631,
        "heading": "4.1.2.1 หน้าภาพรวมระบบ",
        "image": "BackofficeDashboard.png",
        "caption": "รูปที่ 4.13 หน้าภาพรวมระบบ Backoffice",
        "description": (
            "ภาพนี้แสดงหน้าภาพรวมสำหรับผู้ดูแลระบบ ซึ่งสรุปจำนวนผู้ใช้ ทริป "
            "คำขอยืนยันตัวตน รายงานที่รอตรวจสอบ บัญชีที่ยืนยันแล้ว และบัญชีที่ถูกระงับ "
            "พร้อมส่วนงานที่ควรตรวจสอบเพื่อให้ผู้ดูแลเข้าถึงรายการสำคัญได้รวดเร็ว"
        ),
    },
    {
        "paragraph": 635,
        "image_paragraph": 636,
        "caption_paragraph": 637,
        "heading": "4.1.2.2 หน้าจัดการผู้ใช้",
        "image": "BackofficeUsers.png",
        "caption": "รูปที่ 4.14 หน้าจัดการผู้ใช้",
        "description": (
            "ภาพนี้แสดงรายชื่อผู้ใช้ในรูปแบบตาราง โดยผู้ดูแลสามารถค้นหาจากชื่อหรืออีเมล "
            "กรองตามสถานะ และตรวจสอบจำนวนทริปที่สร้าง จำนวนทริปที่เข้าร่วม จำนวนรายงาน "
            "รวมถึงวันที่สมัครและสถานะบัญชี ก่อนดำเนินการระงับบัญชีที่ไม่เป็นไปตามข้อกำหนด"
        ),
    },
    {
        "paragraph": 644,
        "image_paragraph": 645,
        "caption_paragraph": 646,
        "heading": "4.1.2.3 หน้าจัดการทริป",
        "image": "BackofficeTrips.png",
        "caption": "รูปที่ 4.15 หน้าจัดการทริป",
        "description": (
            "ภาพนี้แสดงรายการทริปทั้งหมดพร้อมชื่อทริป จังหวัด ผู้สร้าง จำนวนสมาชิก "
            "วันเริ่มเดินทาง งบประมาณ และหมวดหมู่ ผู้ดูแลสามารถค้นหาทริป เปิดดูรายละเอียด "
            "หรือลบทริปที่มีข้อมูลไม่เหมาะสมและขัดต่อข้อกำหนดของระบบได้"
        ),
    },
    {
        "paragraph": 651,
        "image_paragraph": 652,
        "caption_paragraph": 653,
        "heading": "4.1.2.4 หน้าจัดการรายงานผู้ใช้",
        "image": "BackofficeReport.png",
        "caption": "รูปที่ 4.16 หน้าจัดการรายงานผู้ใช้",
        "description": (
            "ภาพนี้แสดงส่วนตรวจสอบรายงานบัญชีที่ผู้ใช้ส่งเข้ามา เมื่อมีรายงาน ระบบจะแสดงผู้รายงาน "
            "ผู้ถูกรายงาน เหตุผล และสถานะของรายการ เพื่อให้ผู้ดูแลพิจารณาหลักฐาน "
            "ส่งคำเตือน ระงับบัญชี หรือปิดรายงานหลังดำเนินการเสร็จสิ้น"
        ),
    },
    {
        "paragraph": 658,
        "image_paragraph": 659,
        "caption_paragraph": 660,
        "heading": "4.1.2.5 หน้าตรวจสอบการยืนยันตัวตน",
        "image": "BackofficeVerified.png",
        "caption": "รูปที่ 4.17 หน้าตรวจสอบการยืนยันตัวตน",
        "description": (
            "ภาพนี้แสดงรายการคำขอยืนยันตัวตนที่รอการตรวจสอบ ผู้ดูแลใช้หน้านี้ตรวจสอบข้อมูลและภาพใบหน้า "
            "ที่ผู้ใช้ส่งจากแอป ก่อนเลือกอนุมัติหรือปฏิเสธคำขอ ผลการพิจารณาจะนำไปอัปเดต "
            "สถานะการยืนยันตัวตนของบัญชีผู้ใช้"
        ),
    },
    {
        "paragraph": 665,
        "image_paragraph": 666,
        "caption_paragraph": 667,
        "heading": "4.1.2.6 หน้าจัดการและประกาศการแจ้งเตือน",
        "image": "BackofficeNoti.png",
        "caption": "รูปที่ 4.18 หน้าจัดการและประกาศการแจ้งเตือนระบบ",
        "description": (
            "ภาพนี้แสดงแบบฟอร์มสร้างการแจ้งเตือนและรายการแจ้งเตือนที่ระบบเคยส่ง "
            "ผู้ดูแลสามารถระบุหัวข้อ ข้อความ และประเภทการแจ้งเตือน ได้แก่ ทั่วไป ทริป หรือระบบ "
            "จากนั้นส่งประกาศให้ผู้ใช้ รวมถึงตรวจสอบและล้างรายการแจ้งเตือนได้"
        ),
    },
]


def set_run_font(run, size=16, bold=False):
    run.font.name = "TH Sarabun New"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = None
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), "TH Sarabun New")


def replace_paragraph_text(paragraph, text, style=None, size=16, bold=False):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    if style:
        paragraph.style = style
    # Remove direct list numbering retained from the source paragraph.  Changing
    # the style alone is not enough when the original paragraph has a w:numPr.
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def insert_paragraph_after(paragraph, text, style="Normal (Web)"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    new_p.addnext(new_paragraph._p)
    new_paragraph.style = style
    run = new_paragraph.add_run(text)
    set_run_font(run, size=16)
    new_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    new_paragraph.paragraph_format.first_line_indent = Inches(0.5)
    new_paragraph.paragraph_format.space_after = Pt(8)
    return new_paragraph


def replace_image(paragraph, image_path, width):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=width)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    original_paragraphs = list(document.paragraphs)
    section = document.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin

    for item in SECTIONS:
        image_path = IMAGE_DIR / item["image"]
        if not image_path.exists():
            raise FileNotFoundError(image_path)

        heading_p = original_paragraphs[item["paragraph"]]
        replace_paragraph_text(
            heading_p,
            item["heading"],
            # The source document's Heading 3 style has automatic list numbering.
            # Use its body style with direct heading formatting so the explicit
            # 4.1.2.x number is not duplicated as "1. 4.1.2.x".
            style="Normal (Web)",
            size=18,
            bold=True,
        )
        heading_p.paragraph_format.keep_with_next = True
        heading_p.paragraph_format.space_before = Pt(12)
        heading_p.paragraph_format.space_after = Pt(5)
        insert_paragraph_after(heading_p, item["description"])

        image_p = original_paragraphs[item["image_paragraph"]]
        replace_image(image_p, image_path, usable_width)

        caption_p = original_paragraphs[item["caption_paragraph"]]
        replace_paragraph_text(caption_p, item["caption"], style="Normal (Web)", size=14)
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_p.paragraph_format.keep_with_previous = True
        caption_p.paragraph_format.space_before = Pt(3)
        caption_p.paragraph_format.space_after = Pt(10)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
