from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE=Path('/Users/worakanp/Downloads/รายงานฉบับเต็ม_แยกคะแนนหมวดหมู่และช่วงเวลา.docx')
OUTPUT=Path('/tmp/report_polished/รายงานฉบับเต็ม_จัดหน้าพร้อมส่ง.docx')
d=Document(SOURCE)
original_images=len(d.inline_shapes)
for s in d.sections:
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2.54);s.bottom_margin=Cm(2.54)
    s.left_margin=Cm(3);s.right_margin=Cm(2.54)
for style in d.styles:
    if hasattr(style,'font'):
        style.font.name='TH Sarabun New';style.font.size=Pt(16)

def font(p):
    for r in p.runs:
        r.font.name='TH Sarabun New';r.font.size=Pt(16)
        rf=r._element.get_or_add_rPr().get_or_add_rFonts()
        for name in ('ascii','hAnsi','eastAsia','cs'):
            rf.set(qn('w:'+name),'TH Sarabun New')

ps=list(d.paragraphs)
for i,p in enumerate(ps):
    font(p)
    if i<36: continue
    if not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:pict|.//w:sectPr|.//w:br|.//m:oMath|.//w:object'):
        p._p.getparent().remove(p._p)
        continue
    f=p.paragraph_format
    f.widow_control=True
    f.space_before=Pt(0);f.space_after=Pt(4)
    f.line_spacing=1.05
    t=p.text.strip()
    isheading=p.style.name.startswith('Heading') or t in ('สารบัญ','บทคัดย่อ','อ้างอิง') or (len(t)<150 and re.match(r'^\d+\.\d+(?:\.\d+)*\s',t) and not p.style.name.startswith('toc'))
    # Existing heading hierarchy controls scope; do not change chapter contents.
    if isheading:
        f.keep_with_next=True;f.keep_together=True
        f.space_before=Pt(8);f.space_after=Pt(6)
        for r in p.runs:r.bold=True
    if re.fullmatch(r'บทที่\s*[1-5]',t) or t in ('สารบัญ','บทคัดย่อ','อ้างอิง'):
        f.page_break_before=True
        f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        f.space_before=Pt(0)
    if p.style.name.startswith('toc'):
        f.space_after=Pt(2);f.line_spacing=1
        f.keep_with_next=False
        if t.startswith('บทที่ 3 '):
            f.page_break_before=True
    if t.startswith(('รูปที่','ภาพที่','คำบรรยาย:','ตารางที่')):
        f.keep_together=True;f.space_before=Pt(3);f.space_after=Pt(8)
        if t.startswith(('รูปที่','ภาพที่','คำบรรยาย:')):
            f.alignment=WD_ALIGN_PARAGRAPH.CENTER
            f.keep_with_next=False
        else:f.keep_with_next=True
    if p._p.xpath('.//w:drawing'):
        f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        f.keep_together=True;f.keep_with_next=True
        f.space_before=Pt(6);f.space_after=Pt(3)

# Scale images proportionally to leave room for a readable caption.
maxw=Cm(15.46);maxh=Cm(18)
for shape in d.inline_shapes:
    scale=min(1,maxw/shape.width,maxh/shape.height)
    if scale<1:
        shape.width=int(shape.width*scale);shape.height=int(shape.height*scale)

for table in d.tables:
    table.autofit=True
    for row in table.rows:
        row.height=None
        pr=row._tr.get_or_add_trPr()
        if not pr.find(qn('w:cantSplit')) is not None:
            pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                font(p)
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.widow_control=True
    if table.rows:
        pr=table.rows[0]._tr.get_or_add_trPr()
        if pr.find(qn('w:tblHeader')) is None:
            pr.append(OxmlElement('w:tblHeader'))

settings=d.settings.element
u=settings.find(qn('w:updateFields'))
if u is None:u=OxmlElement('w:updateFields');settings.append(u)
u.set(qn('w:val'),'true')
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUTPUT)
assert len(Document(OUTPUT).inline_shapes)==original_images
print(OUTPUT)
