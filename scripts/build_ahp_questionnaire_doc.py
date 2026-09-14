from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "GoWithUs_แบบสอบถามหาน้ำหนัก_AHP.docx"


def font(run, size=14, bold=False, color=None):
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    prop = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    prop.append(node)


def text(cell, value, *, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(value)
    font(r, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def heading(value, size=17):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(value)
    font(r, size=size, bold=True, color="24449A")
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)

style = doc.styles["Normal"]
style.font.name = "TH Sarabun New"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
style.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("แบบสอบถามความสำคัญของปัจจัยในการเลือกทริป")
font(r, size=20, bold=True, color="24449A")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("สำหรับกำหนดน้ำหนักระบบ Smart Matching ของแอปพลิเคชัน GoWithUs")
font(r, size=16, bold=True)

heading("คำชี้แจง")
doc.add_paragraph(
    "แบบสอบถามนี้มีวัตถุประสงค์เพื่อศึกษาความสำคัญของปัจจัยที่ผู้ใช้พิจารณาเมื่อตัดสินใจเลือกทริป "
    "คำตอบจะนำไปคำนวณค่าน้ำหนักด้วยกระบวนการ Analytic Hierarchy Process (AHP) "
    "การตอบเป็นความสมัครใจ ไม่มีคำตอบถูกหรือผิด และรายงานผลในภาพรวมโดยไม่ระบุตัวตน"
)

heading("ความยินยอมและคุณสมบัติผู้ตอบ")
doc.add_paragraph("1. ท่านยินยอมเข้าร่วมตอบแบบสอบถามโดยสมัครใจหรือไม่?")
doc.add_paragraph("☐ ยินยอม     ☐ ไม่ยินยอม (ยุติการตอบแบบสอบถาม)")
doc.add_paragraph("2. ท่านเคยท่องเที่ยวร่วมกับผู้อื่น หรือสนใจใช้แอปพลิเคชันหาเพื่อนร่วมทริปหรือไม่?")
doc.add_paragraph("☐ เคยหรือสนใจ     ☐ ไม่เคยและไม่สนใจ (ยุติการตอบแบบสอบถาม)")

heading("ความหมายของปัจจัย")
definitions = [
    ("ความสนใจ", "ประเภทการท่องเที่ยวตรงกับสิ่งที่ชอบ เช่น ทะเล ภูเขา คาเฟ่ หรือไหว้พระ"),
    ("งบประมาณ", "ค่าใช้จ่ายเฉลี่ยต่อคนของทริปอยู่ในระดับที่สามารถจ่ายได้"),
    ("จำนวนกิจกรรมต่อวัน", "ความหนาแน่นของตารางกิจกรรมตรงกับรูปแบบการเที่ยวที่ต้องการ"),
    ("ช่วงเวลา", "ช่วงเช้า กลางวัน เย็น หรือกลางคืนตรงกับช่วงที่สะดวกหรือชอบทำกิจกรรม"),
]
t = doc.add_table(rows=1, cols=2)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(("ปัจจัย", "ความหมาย")):
    text(t.rows[0].cells[i], h, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(t.rows[0].cells[i], "24449A")
for name, meaning in definitions:
    cells = t.add_row().cells
    text(cells[0], name, bold=True)
    text(cells[1], meaning)

heading("วิธีตอบ")
doc.add_paragraph(
    "แต่ละข้อให้เปรียบเทียบปัจจัย 2 ด้าน แล้วเลือกเพียงหนึ่งคำตอบ เลข 1–9 คือระดับความสำคัญ "
    "ไม่ใช่เปอร์เซ็นต์: 1 = สำคัญเท่ากัน, 3 = สำคัญกว่าปานกลาง, 5 = สำคัญกว่ามาก, "
    "7 = สำคัญกว่าอย่างชัดเจน และ 9 = สำคัญกว่าอย่างยิ่ง ส่วน 2, 4, 6 และ 8 เป็นค่าระหว่างระดับ"
)

pairs = [
    ("งบประมาณ", "จำนวนกิจกรรมต่อวัน"),
    ("งบประมาณ", "ความสนใจ"),
    ("งบประมาณ", "ช่วงเวลา"),
    ("จำนวนกิจกรรมต่อวัน", "ความสนใจ"),
    ("จำนวนกิจกรรมต่อวัน", "ช่วงเวลา"),
    ("ความสนใจ", "ช่วงเวลา"),
]

heading("คำถามเปรียบเทียบปัจจัย")
for idx, (a, b) in enumerate(pairs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    r = p.add_run(f"ข้อ {idx} ระหว่าง “{a}” กับ “{b}” ปัจจัยใดสำคัญกว่าต่อการเลือกทริปของท่าน?")
    font(r, bold=True)
    options = [
        f"☐ {a} สำคัญกว่า {b} ระดับ 9", f"☐ {a} สำคัญกว่า {b} ระดับ 8",
        f"☐ {a} สำคัญกว่า {b} ระดับ 7", f"☐ {a} สำคัญกว่า {b} ระดับ 6",
        f"☐ {a} สำคัญกว่า {b} ระดับ 5", f"☐ {a} สำคัญกว่า {b} ระดับ 4",
        f"☐ {a} สำคัญกว่า {b} ระดับ 3", f"☐ {a} สำคัญกว่า {b} ระดับ 2",
        "☐ ทั้งสองปัจจัยสำคัญเท่ากัน (ระดับ 1)",
        f"☐ {b} สำคัญกว่า {a} ระดับ 2", f"☐ {b} สำคัญกว่า {a} ระดับ 3",
        f"☐ {b} สำคัญกว่า {a} ระดับ 4", f"☐ {b} สำคัญกว่า {a} ระดับ 5",
        f"☐ {b} สำคัญกว่า {a} ระดับ 6", f"☐ {b} สำคัญกว่า {a} ระดับ 7",
        f"☐ {b} สำคัญกว่า {a} ระดับ 8", f"☐ {b} สำคัญกว่า {a} ระดับ 9",
    ]
    grid = doc.add_table(rows=9, cols=2)
    grid.style = "Table Grid"
    for n, option in enumerate(options):
        row, col = divmod(n, 2)
        text(grid.rows[row].cells[col], option)
    if len(options) % 2:
        text(grid.rows[-1].cells[-1], "")

heading("ข้อเสนอแนะเพิ่มเติม (ไม่บังคับ)")
doc.add_paragraph("มีปัจจัยอื่นที่สำคัญต่อการเลือกทริป หรือมีข้อใดที่เข้าใจยากหรือไม่?")
doc.add_paragraph("............................................................................................................................")
doc.add_paragraph("............................................................................................................................")

heading("หมายเหตุสำหรับผู้วิจัย (ไม่นำไปแสดงแก่ผู้ตอบ)")
doc.add_paragraph(
    "แบบสอบถามนี้มีคำถามเปรียบเทียบครบ 6 คู่สำหรับปัจจัย 4 ด้าน ห้ามแสดงน้ำหนักเดิม "
    "35/30/20/15 แก่ผู้ตอบ และห้ามปรับคำตอบเพื่อให้ได้ผลตามน้ำหนักเดิม หลังเก็บข้อมูลให้คำนวณ "
    "AHP พร้อมตรวจ Consistency Ratio และรายงานจำนวนผู้ตอบ วิธีคัดเลือกกลุ่มตัวอย่าง และวิธีรวมคำตอบอย่างชัดเจน"
)
doc.add_paragraph(
    "อ้างอิงวิธีการ: Saaty, T. L. (2008). Decision making with the analytic hierarchy process. "
    "International Journal of Services Sciences, 1(1), 83–98. https://doi.org/10.1504/IJSSCI.2008.017590"
)

doc.save(OUTPUT)
print(OUTPUT)
