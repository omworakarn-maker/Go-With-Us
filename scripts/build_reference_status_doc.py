from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "GoWithUs_ตารางสถานะแหล่งอ้างอิง.docx"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(14)
    if color:
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles[f"Heading {level}"]
    run = paragraph.add_run(text)
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

normal = doc.styles["Normal"]
normal.font.name = "TH Sarabun New"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
normal.font.size = Pt(14)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("ตารางตรวจสอบสถานะแหล่งอ้างอิงของระบบ Smart Matching")
run.bold = True
run.font.name = "TH Sarabun New"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
run.font.size = Pt(20)

intro = doc.add_paragraph(
    "เอกสารนี้แยกระหว่าง (1) แนวคิดหรือวิธีคำนวณที่มีเอกสารวิชาการรองรับ "
    "และ (2) ค่าพารามิเตอร์เฉพาะของโครงการซึ่งยังต้องเก็บข้อมูลหรือประเมินเพิ่มเติม "
    "การมีแหล่งอ้างอิงรองรับสูตรไม่ได้หมายความว่าแหล่งนั้นรับรองตัวเลขทุกค่าที่กำหนดในระบบ"
)
intro.paragraph_format.space_after = Pt(8)

rows = [
    ("Cosine Similarity และ Dot Product", "มี", "Manning, Raghavan และ Schütze (2008)", "รองรับสูตรวัดความคล้ายของเวกเตอร์ แต่ไม่ได้รับรองน้ำหนักหรือเกณฑ์ของ GoWithUs"),
    ("การเข้ารหัสแบบ 0/1 และ Multi-hot", "มี", "แนวคิด Vector Space Model และการเข้ารหัสคุณลักษณะเชิงหมวดหมู่", "รองรับการแทนการเลือก/ไม่เลือกด้วยตัวเลข แต่จำนวนหมวดเป็นการออกแบบของโครงการ"),
    ("การทำ Unit Vector หรือ Normalization", "มี", "Manning, Raghavan และ Schütze (2008); OECD/EC-JRC (2008)", "รองรับหลักการปรับสเกลและการทำเวกเตอร์หน่วย"),
    ("การใช้หลายปัจจัยในระบบแนะนำการท่องเที่ยว", "มี", "Zhou, Tian, Peng และ Su (2021)", "รองรับการพิจารณาปัจจัยหลายด้าน เช่น งบประมาณและเวลา แต่ไม่ได้กำหนดน้ำหนักชุดเดียวกับโครงการ"),
    ("วิธีหาน้ำหนักด้วย AHP", "มี", "Saaty (2008); Zaman, Botti และ Vo-Thanh (2016)", "รองรับกระบวนการหาน้ำหนักจากการเปรียบเทียบปัจจัย ไม่ได้กำหนดผลลัพธ์ 35/30/20/15"),
    ("น้ำหนัก 35/30/20/15", "ยังไม่มีโดยตรง", "เป็นค่าคงที่ที่ผู้วิจัย/ผู้พัฒนากำหนด", "ต้องมีแบบสอบถาม AHP ข้อมูลผู้ประเมิน วิธีคำนวณ และผลตรวจความสอดคล้อง"),
    ("การแบ่งระดับกิจกรรมเป็น 2/5/8", "ยังไม่มีโดยตรง", "เป็นรหัสระดับที่โครงการกำหนด", "ต้องอธิบายเกณฑ์แบ่งกลุ่มและทดสอบกับข้อมูลผู้ใช้จริง"),
    ("การหารระดับกิจกรรมด้วย 6", "ยังไม่มีโดยตรง", "มาจากช่วง 8−2 ของสเกลในโค้ด", "เป็นวิธี Normalize ที่ผูกกับสเกลของโครงการ จึงต้องมีเหตุผลและการประเมินผล"),
    ("จุดตัดงบประมาณ 50%", "ยังไม่มีโดยตรง", "เป็นกฎของโครงการ", "ต้องเก็บข้อมูลการยอมรับทริปเมื่องบประมาณเกินกว่าที่ผู้ใช้กำหนด"),
    ("เพดานคะแนน 39% เมื่อเกินงบสองเท่า", "ยังไม่มีโดยตรง", "เป็นกฎของโครงการ", "ต้องเปรียบเทียบกับค่าเพดานอื่นและประเมินผลด้วยข้อมูลจริง"),
    ("ช่วงแปลผล 25/50/75", "ยังไม่มีโดยตรง", "เป็นเกณฑ์การแสดงผลของ UI", "ต้องตรวจว่าช่วงคะแนนสอดคล้องกับความเห็นหรือพฤติกรรมของผู้ใช้จริง"),
    ("จำนวนหมวดความสนใจและจำนวนที่เลือกได้", "ยังไม่มีโดยตรง", "เป็นขอบเขตที่โครงการกำหนด", "ต้องแสดงที่มาของหมวดและทดสอบความครอบคลุมกับกลุ่มเป้าหมาย"),
]

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["ประเด็น", "สถานะ", "ที่มา/แหล่งอ้างอิง", "ข้อสรุปและหลักฐานที่ต้องเพิ่ม"]
for i, header in enumerate(headers):
    set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF")
    set_cell_shading(table.rows[0].cells[i], "24449A")

for issue, status, source, note in rows:
    cells = table.add_row().cells
    set_cell_text(cells[0], issue)
    set_cell_text(cells[1], status, bold=True, color="16803A" if status == "มี" else "B54708")
    set_cell_text(cells[2], source)
    set_cell_text(cells[3], note)

widths = [Cm(4.3), Cm(2.6), Cm(5.0), Cm(6.0)]
for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = width

add_heading(doc, "ข้อสรุปสำคัญ", 1)
for text in [
    "มีแหล่งอ้างอิงรองรับการใช้ Cosine Similarity, Dot Product, Normalization, การแนะนำแบบหลายปัจจัย และวิธี AHP",
    "ยังไม่พบงานวิจัยที่กำหนดน้ำหนัก 35%, 30%, 20% และ 15% ตรงกับระบบนี้",
    "ค่าระดับ 2/5/8 จุดตัดงบ 50% เพดาน 39% และช่วงแปลผล 25/50/75 เป็นค่าที่โครงการกำหนดและยังต้องยืนยันด้วยข้อมูลจริง",
    "หากต้องการอ้างตัวเลขน้ำหนักเป็นผลวิจัย ควรเก็บแบบสอบถามจากกลุ่มเป้าหมายและคำนวณด้วย AHP โดยไม่กำหนดคำตอบล่วงหน้า",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)

add_heading(doc, "รายการอ้างอิง", 1)
references = [
    "[1] Manning, C. D., Raghavan, P., & Schütze, H. (2008). Introduction to Information Retrieval. Cambridge University Press. https://nlp.stanford.edu/IR-book/",
    "[2] Saaty, T. L. (2008). Decision making with the analytic hierarchy process. International Journal of Services Sciences, 1(1), 83–98. https://doi.org/10.1504/IJSSCI.2008.017590",
    "[3] Zhou, X., Tian, J., Peng, J., & Su, M. (2021). A Smart Tourism Recommendation Algorithm Based on Cellular Geospatial Clustering and Multivariate Weighted Collaborative Filtering. ISPRS International Journal of Geo-Information, 10(9), 628. https://doi.org/10.3390/ijgi10090628",
    "[4] Zaman, M., Botti, L., & Vo-Thanh, T. (2016). Weight of criteria in hotel selection: An empirical illustration based on TripAdvisor criteria. European Journal of Tourism Research, 13, 132–138.",
    "[5] OECD, European Union, & Joint Research Centre. (2008). Handbook on Constructing Composite Indicators: Methodology and User Guide. OECD Publishing. https://doi.org/10.1787/9789264043466-en",
]
for reference in references:
    doc.add_paragraph(reference)

doc.save(OUTPUT)
print(OUTPUT)
